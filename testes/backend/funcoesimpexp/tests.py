# testes/backend/funcoesimpexp/tests.py
"""IE.0 — testes do motor.

Concentrados nas **travas da régua** (planta §2), não no caminho feliz: o valor
deste app não é gerar um XLSX, é impedir que um arquivo torto entre no banco sem
alguém olhar (R1), que reimportar duplique (R3/R14), que meia permissão exporte
metade de uma planilha (R11) e que um lote ruim vire trabalho perdido (R6).

O recurso de teste grava num dicionário em memória em vez de num modelo real: o
que está sob teste é o motor, e amarrá-lo a um domínio faria o teste falhar por
motivo alheio no dia em que aquele domínio mudasse.
"""
from __future__ import annotations

import tempfile
from datetime import timedelta
from decimal import Decimal

from django.core.files.base import ContentFile
from django.test import TestCase, override_settings
from django.utils import timezone
from rest_framework.exceptions import PermissionDenied
from rest_framework.test import APITestCase

from nucleo.models import Tenant, Usuario

from funcoesimpexp import exportacao as motor_exportacao
from funcoesimpexp import importacao as motor_importacao
from funcoesimpexp.acesso import permissoes_faltantes, pode_operar
from funcoesimpexp.formatos.leitura import detectar_delimitador, para_decimal
from funcoesimpexp.models import Exportacao, Importacao, LinhaImportacao, MapaColunas
from funcoesimpexp.modelo import planilha_modelo
from funcoesimpexp.registro import (
    Coluna, Recurso, RecursoInvalido, _REGISTRO, listar, obter, registrar,
)
from funcoesimpexp.tasks import gerar_exportacao, limpar_exportacoes_expiradas, preparar_importacao

PERMS = ("teste.item.view", "teste.item.edit", "teste.item_doc.view")


# ---------------------------------------------------------------------------
# Troca do registry — apoio só deste arquivo
# ---------------------------------------------------------------------------
# O `registro.py` é produção: declara o contrato (`Recurso`, `Coluna`) e o
# registry povoado no `ready()` do app. As três funções abaixo só existem para um
# teste guardar o estado do registry, trocar pelos recursos de teste e devolver
# ao que era — por isso moram aqui, e não lá.
#
# Acessam `_REGISTRO` (privado de `registro`) de propósito: o ponto de `limpar` é
# esvaziar exatamente o dict que o `ready()` povoou, e o de `restaurar` é
# repreenchê-lo. Importar a referência ao dict mutável funciona porque `clear` e
# `update` agem sobre o objeto original — não há cópia.


def instantaneo() -> dict[str, Recurso]:
    """Cópia do registry, para um teste devolvê-lo ao que era (ver `limpar`)."""
    return dict(_REGISTRO)


def restaurar(estado: dict[str, Recurso]) -> None:
    _REGISTRO.clear()
    _REGISTRO.update(estado)


def limpar() -> None:
    """Esvazia o registry. **Sempre em par com `instantaneo`/`restaurar`.**

    O registry é global e povoado no `ready()` do app: quem chama `limpar()` sem
    guardar o estado antes apaga as declarações **reais** de todos os módulos, e
    o estrago não aparece na suíte daquele app — aparece nos testes que rodarem
    depois, como recurso que "não existe" (404). Foi o que aconteceu ao rodar a
    suíte inteira pela primeira vez com o Orçamento declarado.
    """
    _REGISTRO.clear()


# ---------------------------------------------------------------------------
# Recurso de teste — destino em memória
# ---------------------------------------------------------------------------

BANCO: dict[str, dict] = {}


class ItemDeTeste(Recurso):
    rotulo = "Itens de teste"
    modulo = "teste"
    permissoes_exigidas = PERMS
    chave = "codigo"
    colunas = (
        Coluna("codigo", "Código", obrigatoria=True, exemplo="A-1"),
        Coluna("descricao", "Descrição", obrigatoria=True, aliases=("nome",)),
        Coluna("quantidade", "Quantidade", tipo=Decimal, exemplo="1,50"),
    )

    def queryset(self, request, recorte):
        itens = list(BANCO.values())
        if recorte.get("codigo"):
            itens = [i for i in itens if i["codigo"] == recorte["codigo"]]
        return itens

    def validar_linha(self, dados):
        if dados["quantidade"] < 0:
            return "Quantidade negativa."
        return ""

    def localizar(self, dados, tenant_id):
        return BANCO.get(dados["codigo"])

    def aplicar_linha(self, dados, existente, tenant_id):
        if existente is None:
            BANCO[dados["codigo"]] = dict(dados)
            return {"acao": "criou", "chave": dados["codigo"]}
        antes = dict(existente)
        BANCO[dados["codigo"]] = dict(dados)
        return {"acao": "atualizou", "chave": dados["codigo"], "antes": _texto(antes)}

    def desfazer_linha(self, estado_anterior, tenant_id):
        if estado_anterior["acao"] == "criou":
            BANCO.pop(estado_anterior["chave"], None)
        else:
            restaurado = dict(estado_anterior["antes"])
            restaurado["quantidade"] = para_decimal(restaurado["quantidade"])
            BANCO[estado_anterior["chave"]] = restaurado


def _texto(dados: dict) -> dict:
    return {k: (str(v) if isinstance(v, Decimal) else v) for k, v in dados.items()}


CABECALHO = "Código;Descrição;Quantidade"


def _csv(linhas: list[str]) -> bytes:
    return "\n".join(linhas).encode("utf-8")


class MotorTestBase(TestCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        # Guarda o que o `ready()` registrou: `limpar()` apaga as declarações
        # REAIS de todos os módulos, e sem restaurar o estrago só aparece nos
        # testes que rodarem depois, como recurso inexistente.
        cls._registro_do_boot = instantaneo()
        limpar()
        cls.recurso = registrar("teste.item")(ItemDeTeste)()

    @classmethod
    def tearDownClass(cls):
        restaurar(cls._registro_do_boot)
        super().tearDownClass()

    @classmethod
    def setUpTestData(cls):
        cls.tenant = Tenant.objects.create(slug="ie-a", nome="Construtora A")
        cls.tenant_b = Tenant.objects.create(slug="ie-b", nome="Construtora B")
        cls.user = Usuario.objects.create(username="operador_a", tenant=cls.tenant)

    def setUp(self):
        super().setUp()
        BANCO.clear()

    def _preparar(self, linhas, **kwargs):
        return motor_importacao.preparar(
            recurso=self.recurso, conteudo=_csv(linhas), nome_arquivo="itens.csv",
            tenant_id=self.tenant.id, usuario_nome="operador", **kwargs,
        )


# ---------------------------------------------------------------------------
# O contrato
# ---------------------------------------------------------------------------


class ContratoTests(TestCase):
    """Declaração torta falha no registro, não no clique do usuário."""

    def setUp(self):
        super().setUp()
        self._registro_do_boot = instantaneo()

    def tearDown(self):
        restaurar(self._registro_do_boot)

    def test_recurso_sem_permissoes_nao_registra(self):
        class SemGate(Recurso):
            rotulo, modulo, chave = "X", "teste", "codigo"
            colunas = (Coluna("codigo", "Código"),)

            def queryset(self, request, recorte):
                return []

        with self.assertRaisesMessage(RecursoInvalido, "permissoes_exigidas"):
            registrar("teste.sem_gate")(SemGate)

    def test_importacao_sem_chave_nao_registra(self):
        class SemChave(Recurso):
            rotulo, modulo = "X", "teste"
            permissoes_exigidas = ("teste.x.view",)
            colunas = (Coluna("codigo", "Código"),)

            def localizar(self, dados, tenant_id):
                return None

            def aplicar_linha(self, dados, existente, tenant_id):
                return None

        with self.assertRaisesMessage(RecursoInvalido, "sem `chave`"):
            registrar("teste.sem_chave")(SemChave)

    def test_recurso_que_nao_faz_nada_nao_registra(self):
        class Inerte(Recurso):
            rotulo, modulo = "X", "teste"
            permissoes_exigidas = ("teste.x.view",)
            colunas = (Coluna("codigo", "Código"),)

        with self.assertRaisesMessage(RecursoInvalido, "não faz nada"):
            registrar("teste.inerte")(Inerte)

    def test_ganchos_da_subclasse_sobrevivem_a_instanciacao(self):
        """O defeito que mataria o desenho se `Recurso` fosse dataclass."""
        registrar("teste.item")(ItemDeTeste)
        recurso = obter("teste.item")
        self.assertTrue(recurso.exporta)
        self.assertTrue(recurso.importa)
        self.assertTrue(recurso.desfazivel)
        # Sem contar o total: o registry tem também o que os módulos reais
        # declararam no boot, e um número fixo aqui quebraria a cada módulo novo
        # que ganhasse import/export — um teste que pune o crescimento normal.
        self.assertIn("teste.item", [r.id for r in listar()])


# ---------------------------------------------------------------------------
# R11 — acesso total ou nada
# ---------------------------------------------------------------------------


class AcessoTotalTests(MotorTestBase):
    def _com(self, perms):
        user = Usuario(username="x", tenant=self.tenant)
        user.iam_permissions = list(perms)
        return user

    def test_permissao_completa_opera(self):
        self.assertTrue(pode_operar(self._com(PERMS), self.recurso))

    def test_meia_permissao_bloqueia(self):
        user = self._com(PERMS[:2])  # falta a de máscara do documento
        self.assertFalse(pode_operar(user, self.recurso))
        self.assertEqual(permissoes_faltantes(user, self.recurso), ("teste.item_doc.view",))

    def test_permissao_nenhuma_bloqueia(self):
        self.assertFalse(pode_operar(self._com([]), self.recurso))

    def test_erro_nomeia_o_que_falta(self):
        from funcoesimpexp.acesso import exigir_acesso_total

        with self.assertRaises(PermissionDenied) as ctx:
            exigir_acesso_total(self._com(PERMS[:1]), self.recurso)
        self.assertIn("teste.item.edit", str(ctx.exception))

    def test_admin_de_tenant_passa(self):
        user = self._com([])
        user.is_tenant_admin = True
        self.assertTrue(pode_operar(user, self.recurso))


# ---------------------------------------------------------------------------
# R1 — nada grava antes da conferência
# ---------------------------------------------------------------------------


class ConferenciaTests(MotorTestBase):
    def test_preparar_classifica_e_nao_grava(self):
        importacao = self._preparar([
            CABECALHO,
            "A-1;Cimento;10",
            "A-2;Areia;5",
            ";Sem código;1",
            "A-1;Cimento repetido;2",
            "A-3;Negativa;-4",
        ])
        self.assertEqual(importacao.status, Importacao.Status.CONFERINDO)
        self.assertEqual(BANCO, {}, "R1: nada pode ir ao destino antes da decisão")

        resumo = motor_importacao.resumo(importacao)
        self.assertEqual(resumo["novas"], 2)
        self.assertEqual(resumo["invalidas"], 3)

        motivos = dict(
            importacao.linhas.filter(situacao=LinhaImportacao.Situacao.INVALIDA)
            .values_list("numero", "motivo")
        )
        self.assertIn("Sem código", motivos[4])
        self.assertIn("repetida dentro do próprio arquivo", motivos[5])
        self.assertEqual(motivos[6], "Quantidade negativa.")

    def test_linha_invalida_nasce_desmarcada(self):
        importacao = self._preparar([CABECALHO, "A-1;Cimento;10", ";Sem código;1"])
        invalida = importacao.linhas.get(situacao=LinhaImportacao.Situacao.INVALIDA)
        self.assertFalse(invalida.aceita)

    def test_aplicar_grava_so_o_aceito(self):
        importacao = self._preparar([CABECALHO, "A-1;Cimento;10", "A-2;Areia;5"])
        alvo = importacao.linhas.get(chave="A-2")
        motor_importacao.marcar(importacao, aceitar=False, linhas=[alvo.linha_id])
        aplicadas = motor_importacao.aplicar(importacao, recurso=self.recurso)

        self.assertEqual(aplicadas, 1)
        self.assertEqual(sorted(BANCO), ["A-1"])
        self.assertEqual(BANCO["A-1"]["quantidade"], Decimal("10"))

    def test_invalida_marcada_a_mao_ainda_nao_entra(self):
        importacao = self._preparar([CABECALHO, "A-1;Cimento;10", "A-3;Negativa;-4"])
        motor_importacao.marcar(importacao, aceitar=True)  # marca tudo, inclusive a inválida
        motor_importacao.aplicar(importacao, recurso=self.recurso)
        self.assertEqual(sorted(BANCO), ["A-1"])

    def test_descartar_nao_grava_nada(self):
        importacao = self._preparar([CABECALHO, "A-1;Cimento;10"])
        motor_importacao.descartar(importacao)
        self.assertEqual(BANCO, {})
        self.assertEqual(importacao.status, Importacao.Status.DESCARTADA)

    def test_decidida_nao_se_decide_de_novo(self):
        importacao = self._preparar([CABECALHO, "A-1;Cimento;10"])
        motor_importacao.aplicar(importacao, recurso=self.recurso)
        with self.assertRaises(motor_importacao.ImportacaoInvalida):
            motor_importacao.aplicar(importacao, recurso=self.recurso)

    def test_arquivo_ilegivel_vira_registro_com_erro(self):
        """Rastro, não exceção solta: sem isso a pessoa reenvia o mesmo arquivo."""
        importacao = motor_importacao.preparar(
            recurso=self.recurso, conteudo=b"nada", nome_arquivo="foto.png",
            tenant_id=self.tenant.id,
        )
        self.assertEqual(importacao.status, Importacao.Status.ERRO)
        self.assertIn("Formato não suportado", importacao.mensagem)


# ---------------------------------------------------------------------------
# R3/R14 — chave de identidade e idempotência
# ---------------------------------------------------------------------------


class IdentidadeTests(MotorTestBase):
    def test_reimportar_atualiza_em_vez_de_duplicar(self):
        primeira = self._preparar([CABECALHO, "A-1;Cimento;10"])
        motor_importacao.aplicar(primeira, recurso=self.recurso)

        segunda = self._preparar([CABECALHO, "A-1;Cimento CP-II;25"])
        linha = segunda.linhas.get(chave="A-1")
        self.assertEqual(linha.situacao, LinhaImportacao.Situacao.DUPLICADA)

        motor_importacao.aplicar(segunda, recurso=self.recurso)
        self.assertEqual(len(BANCO), 1, "R14: reimportar o mesmo arquivo não duplica")
        self.assertEqual(BANCO["A-1"]["quantidade"], Decimal("25"))

    def test_desfazer_restaura_o_estado_anterior(self):
        primeira = self._preparar([CABECALHO, "A-1;Cimento;10"])
        motor_importacao.aplicar(primeira, recurso=self.recurso)

        segunda = self._preparar([CABECALHO, "A-1;Trocado;99", "A-2;Areia;5"])
        motor_importacao.aplicar(segunda, recurso=self.recurso)
        self.assertEqual(BANCO["A-1"]["quantidade"], Decimal("99"))

        motor_importacao.desfazer(segunda, recurso=self.recurso)
        self.assertEqual(BANCO["A-1"]["quantidade"], Decimal("10"), "R6: valor anterior volta")
        self.assertEqual(BANCO["A-1"]["descricao"], "Cimento", "R6: restaura a linha inteira")
        self.assertNotIn("A-2", BANCO, "R6: desfazer uma criação é apagar")

    def test_so_lote_aplicado_se_desfaz(self):
        importacao = self._preparar([CABECALHO, "A-1;Cimento;10"])
        with self.assertRaises(motor_importacao.ImportacaoInvalida):
            motor_importacao.desfazer(importacao, recurso=self.recurso)


# ---------------------------------------------------------------------------
# R4/R8 — leitura brasileira e de-para
# ---------------------------------------------------------------------------


class LeituraTests(MotorTestBase):
    def test_ponto_e_virgula_vence_a_virgula_decimal(self):
        """`csv.Sniffer` elegeria a vírgula e partiria a última coluna em duas."""
        self.assertEqual(detectar_delimitador("A-1;Cimento;0,58"), ";")

    def test_numero_brasileiro(self):
        self.assertEqual(para_decimal("1.234,56"), Decimal("1234.56"))
        self.assertEqual(para_decimal("1234,56"), Decimal("1234.56"))
        self.assertEqual(para_decimal("R$ 87,30"), Decimal("87.30"))
        self.assertEqual(para_decimal(None), Decimal("0"))

    def test_alias_de_cabecalho_casa(self):
        importacao = self._preparar(["Código;Nome;Quantidade", "A-1;Cimento;10"])
        self.assertEqual(importacao.linhas.get().valores["descricao"], "Cimento")

    def test_cabecalho_irreconhecivel_pede_de_para(self):
        with self.assertRaises(motor_importacao.ColunasNaoReconhecidas) as ctx:
            self._preparar(["Xis;Ipsilon;Ze", "A-1;Cimento;10"])
        self.assertIn("Código", ctx.exception.obrigatorias)
        self.assertTrue(ctx.exception.cabecalhos_vistos)

    def test_de_para_manual_destrava(self):
        importacao = self._preparar(
            ["Xis;Ipsilon;Ze", "A-1;Cimento;10"],
            de_para={"codigo": 0, "descricao": 1, "quantidade": 2},
        )
        linha = importacao.linhas.get()
        self.assertEqual(linha.chave, "A-1")
        self.assertEqual(linha.valores["descricao"], "Cimento")

    def test_cabecalho_depois_de_titulo_e_legenda(self):
        importacao = self._preparar([
            "TABELA DE ITENS", "Publicação de julho", "", CABECALHO, "A-1;Cimento;10",
        ])
        self.assertEqual(importacao.total_linhas, 1)


# ---------------------------------------------------------------------------
# R13 — teto do plano
# ---------------------------------------------------------------------------


class TetoDoPlanoTests(MotorTestBase):
    def test_arquivo_acima_do_teto_e_recusado_antes_de_materializar(self):
        linhas = [CABECALHO] + [f"A-{i};Item {i};1" for i in range(5)]
        with self.assertRaises(motor_importacao.LimiteDoPlano) as ctx:
            self._preparar(linhas, limites={"importacao_linhas": 3})
        self.assertEqual(ctx.exception.quota, "importacao_linhas")
        self.assertEqual(Importacao.objects.count(), 0, "R13: nem o lote chega a existir")

    def test_sem_claim_de_quota_nao_trava(self):
        importacao = self._preparar([CABECALHO, "A-1;Cimento;10"], limites=None)
        self.assertEqual(importacao.total_linhas, 1)


# ---------------------------------------------------------------------------
# R2/R7/R9 — modelo, recorte e procedência
# ---------------------------------------------------------------------------


class SaidaTests(MotorTestBase):
    class _Req:
        def __init__(self, user):
            self.user = user

    def test_modelo_sai_das_mesmas_colunas(self):
        from openpyxl import load_workbook
        import io

        wb = load_workbook(io.BytesIO(planilha_modelo(self.recurso)))
        ws = wb.active
        cabecalhos = [c.value for c in ws[5]]
        self.assertEqual(cabecalhos, ["Código", "Descrição", "Quantidade"])

    def test_exportacao_respeita_o_recorte_da_tela(self):
        BANCO["A-1"] = {"codigo": "A-1", "descricao": "Cimento", "quantidade": Decimal("10")}
        BANCO["A-2"] = {"codigo": "A-2", "descricao": "Areia", "quantidade": Decimal("5")}

        conteudo, tipo, nome = motor_exportacao.exportar(
            recurso=self.recurso, request=self._Req(self.user), formato="csv",
            recorte={"codigo": "A-1"}, tenant_id=self.tenant.id,
        )
        texto = conteudo.decode("utf-8-sig")
        self.assertIn("A-1", texto)
        self.assertNotIn("A-2", texto, "R7: o arquivo é o que a pessoa está vendo")
        self.assertTrue(nome.endswith(".csv"))

    def test_csv_sai_com_decimal_brasileiro(self):
        BANCO["A-1"] = {"codigo": "A-1", "descricao": "Cimento", "quantidade": Decimal("1234.56")}
        conteudo, _, _ = motor_exportacao.exportar(
            recurso=self.recurso, request=self._Req(self.user), formato="csv",
            tenant_id=self.tenant.id,
        )
        self.assertIn("1234,56", conteudo.decode("utf-8-sig"))

    def test_xlsx_grava_numero_como_numero(self):
        from openpyxl import load_workbook
        import io

        BANCO["A-1"] = {"codigo": "A-1", "descricao": "Cimento", "quantidade": Decimal("10.5")}
        conteudo, _, _ = motor_exportacao.exportar(
            recurso=self.recurso, request=self._Req(self.user), formato="xlsx",
            tenant_id=self.tenant.id,
        )
        ws = load_workbook(io.BytesIO(conteudo)).active
        valor = ws.cell(row=ws.max_row, column=3).value
        self.assertIsInstance(valor, float, "R8: número tem de chegar somável no Excel")
        self.assertEqual(valor, 10.5)

    def test_exportacao_deixa_rastro(self):
        from funcoesimpexp.models import Exportacao

        BANCO["A-1"] = {"codigo": "A-1", "descricao": "Cimento", "quantidade": Decimal("1")}
        motor_exportacao.exportar(
            recurso=self.recurso, request=self._Req(self.user), formato="csv",
            recorte={"codigo": "A-1"}, tenant_id=self.tenant.id,
        )
        rastro = Exportacao.objects.get()
        self.assertEqual(rastro.recurso, "teste.item")
        self.assertEqual(rastro.linhas, 1)
        self.assertEqual(rastro.recorte, {"codigo": "A-1"})

    def test_formato_nao_oferecido_e_recusado(self):
        with self.assertRaises(motor_exportacao.ExportacaoInvalida):
            motor_exportacao.exportar(
                recurso=self.recurso, request=self._Req(self.user), formato="docx",
                tenant_id=self.tenant.id,
            )


# ---------------------------------------------------------------------------
# R5 — o erro volta como arquivo
# ---------------------------------------------------------------------------


class RecusadasTests(MotorTestBase):
    def test_recusadas_saem_com_o_motivo(self):
        from openpyxl import load_workbook
        import io

        importacao = self._preparar([CABECALHO, "A-1;Cimento;10", "A-3;Negativa;-4"])
        ws = load_workbook(io.BytesIO(
            motor_importacao.recusadas_xlsx(importacao, self.recurso)
        )).active
        linhas = [[c.value for c in linha] for linha in ws.iter_rows(min_row=5)]
        self.assertEqual(linhas[0][-1], "Motivo da recusa")
        self.assertEqual(linhas[1][0], "A-3")
        self.assertEqual(linhas[1][-1], "Quantidade negativa.")


# ---------------------------------------------------------------------------
# Isolamento
# ---------------------------------------------------------------------------


class IsolamentoTests(MotorTestBase):
    def test_importacao_nasce_no_tenant_de_quem_subiu(self):
        importacao = self._preparar([CABECALHO, "A-1;Cimento;10"])
        self.assertEqual(importacao.tenant_id, self.tenant.id)
        self.assertTrue(
            all(l.tenant_id == self.tenant.id for l in importacao.linhas.all())
        )
        self.assertFalse(
            Importacao.objects.filter(tenant_id=self.tenant_b.id).exists()
        )


# ---------------------------------------------------------------------------
# IE.0b — contexto: o destino da operação dentro do módulo
# ---------------------------------------------------------------------------


class ContextoTests(MotorTestBase):
    """O que o IE.1 exigiu do motor: importar *para* algum lugar do módulo."""

    def test_com_contexto_nao_suja_a_instancia_do_registry(self):
        """O defeito seria invisível em teste e cruzaria destinos em produção."""
        vinculado = self.recurso.com_contexto({"banco": 7})
        self.assertEqual(vinculado.contexto, {"banco": 7})
        self.assertEqual(self.recurso.contexto, {}, "o singleton do registry não muda")

        outro = self.recurso.com_contexto({"banco": 9})
        self.assertEqual(vinculado.contexto, {"banco": 7}, "dois destinos não se cruzam")
        self.assertEqual(outro.contexto, {"banco": 9})

    def test_contexto_fica_gravado_no_lote_e_volta_na_decisao(self):
        vinculado = self.recurso.com_contexto({"banco": 7})
        importacao = motor_importacao.preparar(
            recurso=vinculado, conteudo=_csv([CABECALHO, "A-1;Cimento;10"]),
            nome_arquivo="itens.csv", tenant_id=self.tenant.id,
        )
        # Persistido, e não recalculado: entre subir e confirmar, a tela pode ter
        # mudado de banco selecionado.
        self.assertEqual(importacao.contexto, {"banco": 7})

    def test_quota_de_acervo_do_modulo_e_respeitada(self):
        """`importacao_linhas` é do motor; teto de acervo é sempre do módulo."""

        class ComAcervo(ItemDeTeste):
            def checar_quota_objeto(self, novas, limites, tenant_id):
                if novas > 2:
                    raise motor_importacao.LimiteDoPlano(
                        "Acervo cheio.", limite=2, tentado=novas, quota="itens_teste",
                    )

        recurso = ComAcervo()
        with self.assertRaises(motor_importacao.LimiteDoPlano) as ctx:
            motor_importacao.preparar(
                recurso=recurso,
                conteudo=_csv([CABECALHO, "A-1;a;1", "A-2;b;1", "A-3;c;1"]),
                nome_arquivo="itens.csv", tenant_id=self.tenant.id,
            )
        self.assertEqual(ctx.exception.quota, "itens_teste")
        self.assertEqual(Importacao.objects.count(), 0)


# ---------------------------------------------------------------------------
# IE.0b — R10: pequeno é síncrono, grande é job
# ---------------------------------------------------------------------------


class CorteDeVolumeTests(MotorTestBase):
    def test_arquivo_pequeno_nao_vira_job(self):
        self.assertFalse(motor_importacao.vale_a_pena_job(_csv([CABECALHO, "A-1;x;1"])))

    def test_arquivo_grande_vira_job(self):
        grande = _csv([CABECALHO] + [f"A-{i};Item;1" for i in range(4000)])
        self.assertTrue(motor_importacao.vale_a_pena_job(grande))

    def test_exportacao_grande_vira_job(self):
        for i in range(3):
            BANCO[f"A-{i}"] = {"codigo": f"A-{i}", "descricao": "x", "quantidade": Decimal("1")}
        pedido = SaidaTests._Req(self.user)
        self.assertFalse(motor_exportacao.vale_a_pena_job(self.recurso, pedido, {}))

        original = motor_exportacao.LIMITE_SINCRONO
        motor_exportacao.LIMITE_SINCRONO = 2
        try:
            self.assertTrue(motor_exportacao.vale_a_pena_job(self.recurso, pedido, {}))
        finally:
            motor_exportacao.LIMITE_SINCRONO = original


@override_settings(MEDIA_ROOT=tempfile.mkdtemp(prefix="impexp-teste-"))
class JobDeImportacaoTests(MotorTestBase):
    """A task roda no worker: sem request, sem sessão, sob RLS fail-closed (D8)."""

    def _lote_com_arquivo(self, linhas, contexto=None):
        lote = Importacao.objects.create(
            tenant_id=self.tenant.id, recurso=self.recurso.id,
            arquivo_nome="grande.csv", contexto=contexto or {},
            status=Importacao.Status.LENDO, assincrona=True, usuario_nome="operador",
        )
        lote.arquivo.save("grande.csv", ContentFile(_csv(linhas)), save=True)
        return lote

    def test_task_classifica_sem_request_nenhum(self):
        lote = self._lote_com_arquivo([CABECALHO, "A-1;Cimento;10", "A-2;Areia;5"])

        preparar_importacao(lote.pk, self.tenant.id)

        lote.refresh_from_db()
        self.assertEqual(lote.status, Importacao.Status.CONFERINDO)
        self.assertEqual(lote.total_linhas, 2)
        self.assertEqual(lote.progresso_pct, 100)
        self.assertEqual(BANCO, {}, "R1 vale igual no job: nada gravado antes da decisão")

    def test_arquivo_de_entrada_some_depois_de_lido(self):
        """Dado cru (folha, carteira) não fica no disco depois de virar linha."""
        lote = self._lote_com_arquivo([CABECALHO, "A-1;Cimento;10"])
        preparar_importacao(lote.pk, self.tenant.id)
        lote.refresh_from_db()
        self.assertFalse(lote.arquivo)

    def test_lendo_nao_e_decidivel(self):
        lote = self._lote_com_arquivo([CABECALHO, "A-1;Cimento;10"])
        self.assertFalse(lote.pendente)
        with self.assertRaises(motor_importacao.ImportacaoInvalida):
            motor_importacao.aplicar(lote, recurso=self.recurso)

    def test_teto_do_plano_no_job_vira_mensagem_e_nao_falha_de_task(self):
        lote = self._lote_com_arquivo([CABECALHO, "A-1;a;1", "A-2;b;1", "A-3;c;1"])
        preparar_importacao(lote.pk, self.tenant.id, limites={"importacao_linhas": 1})
        lote.refresh_from_db()
        self.assertEqual(lote.status, Importacao.Status.ERRO)
        self.assertIn("plano contratado", lote.mensagem)

    def test_cabecalho_irreconhecivel_no_job_espera_o_de_para(self):
        """R4 vale no job também — e é lá que ela mais vale.

        Arquivo grande é justamente o de origem mais estranha (o PDF da
        prefeitura). Se o job morresse em "não reconheci as colunas", o de-para
        manual — que existe exatamente para esse caso — seria inalcançável no
        único caminho que precisa dele.
        """
        lote = self._lote_com_arquivo(["Xis;Ipsilon;Ze", "A-1;Cimento;10"])
        preparar_importacao(lote.pk, self.tenant.id)

        lote.refresh_from_db()
        self.assertEqual(lote.status, Importacao.Status.ERRO)
        self.assertTrue(lote.parametros.get("aguarda_de_para"))
        self.assertTrue(lote.parametros.get("cabecalhos_vistos"))
        self.assertTrue(lote.arquivo, "o arquivo fica: remapear relê o mesmo, sem novo upload")
        self.assertIsNone(lote.decidida_em, "interrompido não é decidido")

        resumo = motor_importacao.resumo(lote)
        self.assertTrue(resumo["aguarda_de_para"])
        self.assertIn("Código", resumo["obrigatorias"])

    def test_remapear_relê_o_mesmo_arquivo_e_destrava(self):
        lote = self._lote_com_arquivo(["Xis;Ipsilon;Ze", "A-1;Cimento;10"])
        preparar_importacao(lote.pk, self.tenant.id)

        # É o que a view faz no `remapear`: limpa o pedido e reenfileira com o mapa.
        lote.refresh_from_db()
        lote.status = Importacao.Status.LENDO
        lote.parametros = {}
        lote.save(update_fields=["status", "parametros"])
        preparar_importacao(
            lote.pk, self.tenant.id,
            de_para={"codigo": 0, "descricao": 1, "quantidade": 2},
        )

        lote.refresh_from_db()
        self.assertEqual(lote.status, Importacao.Status.CONFERINDO)
        self.assertEqual(lote.total_linhas, 1)
        self.assertEqual(lote.linhas.get().chave, "A-1")
        self.assertFalse(lote.arquivo, "lido com sucesso, o arquivo cru some")


@override_settings(MEDIA_ROOT=tempfile.mkdtemp(prefix="impexp-saida-"))
class JobDeExportacaoTests(MotorTestBase):
    def setUp(self):
        super().setUp()
        BANCO["A-1"] = {"codigo": "A-1", "descricao": "Cimento", "quantidade": Decimal("2")}

    def test_lote_gera_arquivo_com_validade(self):
        lote = motor_exportacao.abrir_lote(
            recurso=self.recurso, formato="csv", recorte={},
            tenant_id=self.tenant.id, usuario_nome="operador",
        )
        self.assertEqual(lote.status, Exportacao.Status.GERANDO)
        self.assertFalse(lote.baixavel)

        gerar_exportacao(lote.pk, self.tenant.id, self.user.pk)

        lote.refresh_from_db()
        self.assertEqual(lote.status, Exportacao.Status.PRONTA)
        self.assertEqual(lote.linhas, 1)
        self.assertTrue(lote.baixavel)
        self.assertIsNotNone(lote.expira_em, "R10: o link tem validade limitada")
        self.assertIn("A-1", lote.arquivo.read().decode("utf-8-sig"))

    def test_expirada_nao_e_baixavel_e_a_faxina_apaga_o_arquivo(self):
        lote = motor_exportacao.abrir_lote(
            recurso=self.recurso, formato="csv", recorte={},
            tenant_id=self.tenant.id, usuario_nome="operador",
        )
        gerar_exportacao(lote.pk, self.tenant.id, self.user.pk)
        Exportacao.objects.filter(pk=lote.pk).update(
            expira_em=timezone.now() - timedelta(minutes=1),
        )
        lote.refresh_from_db()
        self.assertTrue(lote.expirada)
        self.assertFalse(lote.baixavel)

        self.assertEqual(limpar_exportacoes_expiradas(), 1)
        lote.refresh_from_db()
        self.assertFalse(lote.arquivo, "o arquivo some…")
        self.assertEqual(lote.recurso, "teste.item", "…mas o rastro da R12 fica")


# ---------------------------------------------------------------------------
# IE.0b — R4: o de-para salvável
# ---------------------------------------------------------------------------


class ApiTests(APITestCase):
    """Os endpoints — um conjunto só, servindo qualquer recurso.

    O que se prova aqui é o que o motor não alcança sozinho: que o gate do §3
    fecha na porta da API, que o histórico existe (sem ele o desfazer da R6 é
    inalcançável na tela) e que o de-para salvo vai e volta.
    """

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        # Guarda o que o `ready()` registrou: `limpar()` apaga as declarações
        # REAIS de todos os módulos, e sem restaurar o estrago só aparece nos
        # testes que rodarem depois, como recurso inexistente.
        cls._registro_do_boot = instantaneo()
        limpar()
        cls.recurso = registrar("teste.item")(ItemDeTeste)()

    @classmethod
    def tearDownClass(cls):
        restaurar(cls._registro_do_boot)
        super().tearDownClass()

    def setUp(self):
        super().setUp()
        BANCO.clear()
        self.tenant = Tenant.objects.create(slug="ie-api", nome="Construtora API")
        self.user = Usuario.objects.create(username="op_api", tenant=self.tenant)

    def autenticar(self, perms=PERMS):
        self.user.iam_permissions = list(perms)
        self.user.iam_quotas = None
        self.client.force_authenticate(user=self.user)

    def _subir(self, linhas):
        from django.core.files.uploadedfile import SimpleUploadedFile

        arquivo = SimpleUploadedFile("itens.csv", _csv(linhas), content_type="text/csv")
        return self.client.post(
            "/api/impexp/teste.item/importar/", {"arquivo": arquivo}, format="multipart",
        )

    def test_recursos_lista_so_o_que_a_pessoa_pode_operar(self):
        self.autenticar(PERMS[:2])  # falta a de máscara
        self.assertEqual(self.client.get("/api/impexp/recursos/").data["recursos"], [])

        self.autenticar(PERMS)
        ids = [r["id"] for r in self.client.get("/api/impexp/recursos/").data["recursos"]]
        self.assertEqual(ids, ["teste.item"])

    def test_meia_permissao_recebe_403_que_nomeia_o_que_falta(self):
        self.autenticar(PERMS[:2])
        resposta = self._subir([CABECALHO, "A-1;Cimento;10"])
        self.assertEqual(resposta.status_code, 403)
        self.assertIn("teste.item_doc.view", str(resposta.data))

    def test_catalogo_publica_o_teto_antes_do_envio(self):
        """R15 — descobrir o limite no 402, depois do upload, é o que isto evita."""
        self.autenticar()
        limites = self.client.get("/api/impexp/recursos/").data["limites"]
        self.assertGreater(limites["max_mb"], 0)
        self.assertTrue(limites["formatos"], "a allowlist tem de ser publicada")

    def test_saida_de_arquivo_sai_com_attachment_e_nosniff(self):
        """D69 — o par que a PRA.2 fixou como padrão da casa."""
        self.autenticar()
        for url in (
            "/api/impexp/teste.item/exportar/?formato=csv",
            "/api/impexp/teste.item/modelo/",
        ):
            resposta = self.client.get(url)
            self.assertEqual(resposta["X-Content-Type-Options"], "nosniff", url)
            self.assertIn("attachment", resposta["Content-Disposition"], url)

    def test_executavel_renomeado_nao_entra_pela_api(self):
        """D65 — a porta recusa pelos bytes, antes de gravar qualquer coisa."""
        from django.core.files.uploadedfile import SimpleUploadedFile

        self.autenticar()
        arquivo = SimpleUploadedFile(
            "itens.xlsx", b"MZ\x90\x00\x03\x00\x00\x00" + b"\x00" * 64,
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        resposta = self.client.post(
            "/api/impexp/teste.item/importar/", {"arquivo": arquivo}, format="multipart",
        )
        self.assertEqual(resposta.status_code, 400)
        self.assertEqual(Importacao.objects.count(), 0, "nada pode ter sido gravado")

    def test_ciclo_completo_pela_api(self):
        self.autenticar()
        criada = self._subir([CABECALHO, "A-1;Cimento;10", "A-2;Areia;5"])
        self.assertEqual(criada.status_code, 201)
        self.assertEqual(BANCO, {}, "R1: o upload não grava")
        ident = criada.data["importacao_id"]

        detalhe = self.client.get(f"/api/impexp/importacoes/{ident}/")
        self.assertEqual(len(detalhe.data["linhas"]), 2)

        aplicada = self.client.post(f"/api/impexp/importacoes/{ident}/aplicar/")
        self.assertEqual(aplicada.status_code, 200)
        self.assertEqual(sorted(BANCO), ["A-1", "A-2"])

    def test_historico_marca_so_o_ultimo_lote_como_desfazivel(self):
        """Desfazer um lote do meio deixaria um estado que não é nem antes nem depois."""
        self.autenticar()
        primeiro = self._subir([CABECALHO, "A-1;Cimento;10"]).data["importacao_id"]
        self.client.post(f"/api/impexp/importacoes/{primeiro}/aplicar/")
        segundo = self._subir([CABECALHO, "A-2;Areia;5"]).data["importacao_id"]
        self.client.post(f"/api/impexp/importacoes/{segundo}/aplicar/")

        historico = self.client.get("/api/impexp/teste.item/importacoes/").data["importacoes"]
        por_id = {l["importacao_id"]: l["pode_desfazer"] for l in historico}
        self.assertTrue(por_id[segundo])
        self.assertFalse(por_id[primeiro])

        self.client.post(f"/api/impexp/importacoes/{segundo}/desfazer/")
        self.assertEqual(sorted(BANCO), ["A-1"])

    def test_de_para_salvo_vai_e_volta(self):
        self.autenticar()
        criado = self.client.post(
            "/api/impexp/teste.item/mapas/",
            {"nome": "Prefeitura mensal", "mapa": {"codigo": 0, "descricao": 1, "quantidade": 2}},
            format="json",
        )
        self.assertEqual(criado.status_code, 201)

        listados = self.client.get("/api/impexp/teste.item/mapas/").data["mapas"]
        self.assertEqual(listados[0]["nome"], "Prefeitura mensal")

        # E o mapa salvo destrava o arquivo de cabeçalho irreconhecível.
        from django.core.files.uploadedfile import SimpleUploadedFile

        arquivo = SimpleUploadedFile(
            "torto.csv", _csv(["Xis;Ipsilon;Ze", "A-9;Cal;3"]), content_type="text/csv",
        )
        resposta = self.client.post(
            "/api/impexp/teste.item/importar/",
            {"arquivo": arquivo, "mapa": criado.data["mapa_id"]}, format="multipart",
        )
        self.assertEqual(resposta.status_code, 201)
        self.assertEqual(resposta.data["novas"], 1)

    def test_cabecalho_irreconhecivel_responde_422_com_o_que_viu(self):
        self.autenticar()
        resposta = self._subir(["Xis;Ipsilon;Ze", "A-1;Cimento;10"])
        self.assertEqual(resposta.status_code, 422)
        self.assertEqual(resposta.data["code"], "colunas_nao_reconhecidas")
        self.assertTrue(resposta.data["cabecalhos_vistos"])

    def test_recurso_inexistente_da_404(self):
        self.autenticar()
        self.assertEqual(
            self.client.get("/api/impexp/nada.disso/exportar/").status_code, 404,
        )


class MapaSalvoTests(MotorTestBase):
    def test_mapa_salvo_destrava_o_mesmo_arquivo_torto(self):
        mapa = MapaColunas.objects.create(
            tenant_id=self.tenant.id, recurso=self.recurso.id,
            nome="Prefeitura — mensal",
            mapa={"codigo": 0, "descricao": 1, "quantidade": 2},
        )
        importacao = self._preparar(["Xis;Ipsilon;Ze", "A-1;Cimento;10"], de_para=mapa.mapa)
        linha = importacao.linhas.get()
        self.assertEqual(linha.chave, "A-1")
        self.assertEqual(linha.valores["descricao"], "Cimento")

    def test_mapa_e_unico_por_nome_dentro_do_tenant(self):
        from django.db import IntegrityError

        MapaColunas.objects.create(
            tenant_id=self.tenant.id, recurso=self.recurso.id, nome="Mensal", mapa={"codigo": 0},
        )
        # O mesmo nome no outro tenant convive — é o par (tenant, recurso, nome).
        MapaColunas.objects.create(
            tenant_id=self.tenant_b.id, recurso=self.recurso.id, nome="Mensal", mapa={"codigo": 0},
        )
        with self.assertRaises(IntegrityError):
            MapaColunas.objects.create(
                tenant_id=self.tenant.id, recurso=self.recurso.id, nome="Mensal",
                mapa={"codigo": 1},
            )


# ---------------------------------------------------------------------------
# IE.2 — R15 e R16 (endurecimento de 07/08/2026)
# ---------------------------------------------------------------------------
# As travas desta leva não são de fluxo, são de **porta**: o que entra tem de ser
# reconhecido pelos bytes, o que sai não pode executar, e todo teto tem de vir do
# plano. Cada teste aqui corresponde a um achado da análise (D62–D69) — se um
# deles voltar a passar por engano, é porque a porta reabriu.


class FormatoPorAssinaturaTests(TestCase):
    """D65 — o nome do arquivo não decide nada."""

    def test_reconhece_os_tres_formatos_pelo_conteudo(self):
        from funcoesimpexp.limites import detectar_formato

        self.assertEqual(detectar_formato("x.csv", b"a;b;c\n1;2;3").codigo, "csv")
        self.assertEqual(detectar_formato("x.xlsx", b"PK\x03\x04resto").codigo, "xlsx")
        self.assertEqual(detectar_formato("x.pdf", b"%PDF-1.7 resto").codigo, "pdf")

    def test_executavel_renomeado_para_xlsx_e_recusado(self):
        from funcoesimpexp.limites import FormatoDesconhecido, detectar_formato

        # O caso concreto do achado: antes, isto era GRAVADO em disco no caminho
        # de job e só falhava depois, ao ser parseado.
        with self.assertRaises(FormatoDesconhecido):
            detectar_formato("payload.xlsx", b"MZ\x90\x00\x03\x00\x00\x00")
        with self.assertRaises(FormatoDesconhecido):
            detectar_formato("planilha.csv", b"\x7fELF\x02\x01\x01\x00")

    def test_extensao_so_desempata_entre_formatos_de_mesma_assinatura(self):
        from funcoesimpexp.limites import detectar_formato

        # O desempate acontece, e acerta quando o nome é de um formato-zip:
        self.assertEqual(detectar_formato("x.xlsx", b"PK\x03\x04resto").codigo, "xlsx")
        self.assertEqual(detectar_formato("m.bcfzip", b"PK\x03\x04resto").codigo, "bcf")

        # Assinatura de ZIP com nome que não é de nenhum formato-zip conhecido:
        # cai no primeiro que casou, e NÃO vira "desconhecido" — a assinatura já
        # provou que não é executável renomeado.
        #
        # O invariante é esse, e não QUAL deles: a identidade do primeiro mudou
        # (era `xlsx`, virou `bcf`) quando a PRA.5a registrou os 18 formatos da
        # Prancheta na allowlist da plataforma, e amarrá-la aqui faria este teste
        # quebrar a cada formato-zip novo sem que nada tivesse quebrado de fato.
        # Quem decide o destino não é esta função — é o recurso, declarando o que
        # aceita; aqui só se prova que a porta não confunde zip com executável.
        indefinido = detectar_formato("coisa.bin", b"PK\x03\x04resto")
        self.assertIn(b"PK\x03\x04", indefinido.assinaturas)


class LimiteDoPlanoTests(TestCase):
    """D64 — todo teto vem de `payment/planos.py`, e nunca vira ilimitado."""

    class _U:
        def __init__(self, plano=None):
            self.iam_plano = plano
            self.username = "u"

    def test_plano_menor_tem_teto_menor_em_todos_os_eixos(self):
        from funcoesimpexp.limites import limites_do_plano

        free = limites_do_plano(self._U("free"))
        pro = limites_do_plano(self._U("profissional"))
        self.assertLess(free.max_mb, pro.max_mb)
        self.assertLess(free.memoria_mb, pro.memoria_mb)
        self.assertLess(free.segundos, pro.segundos)

    def test_sem_plano_conhecido_cai_no_teto_de_infra_e_nao_em_ilimitado(self):
        from payment.planos import IMPEXP_TETO_INFRA_MB, IMPEXP_TETO_INFRA_MEMORIA_MB

        from funcoesimpexp.limites import limites_do_plano

        # Fail-open é sobre não travar quem tem token de serviço — não sobre
        # entregar a máquina. É a diferença que o `sob_medida` também respeita.
        for plano in (None, "inexistente", "sob_medida"):
            teto = limites_do_plano(self._U(plano))
            self.assertEqual(teto.max_mb, IMPEXP_TETO_INFRA_MB)
            self.assertEqual(teto.memoria_mb, IMPEXP_TETO_INFRA_MEMORIA_MB)

    def test_formato_nao_contratado_da_402_e_arquivo_grande_tambem(self):
        from funcoesimpexp.limites import (
            ArquivoGrandeDemais,
            FormatoNaoContratado,
            limites_do_plano,
            validar_entrada,
        )

        free = limites_do_plano(self._U("free"))
        with self.assertRaises(FormatoNaoContratado):
            validar_entrada(tamanho=10, nome="a.pdf", conteudo=b"%PDF-1.7", limites=free)
        with self.assertRaises(ArquivoGrandeDemais):
            validar_entrada(
                tamanho=free.max_bytes + 1, nome="a.csv", conteudo=b"a;b", limites=free,
            )

    def test_nenhum_plano_vende_formato_que_ninguem_le(self):
        from payment.planos import PLANOS

        from funcoesimpexp.limites import FORMATOS_POR_CODIGO

        # O mesmo teste que o `RegistryEPlanosTests` da Prancheta faz: knob
        # comercial que promete o que o motor não tem é o cliente contratando o
        # que ninguém sabe ler.
        for slug, cfg in PLANOS.items():
            for codigo in cfg.get("impexp_formatos") or ():
                self.assertIn(codigo, FORMATOS_POR_CODIGO, f"plano {slug}")


class FormulaNaSaidaTests(MotorTestBase):
    """D62 — o arquivo que sai não pode executar nada."""

    def _exportar(self, formato, valor):
        BANCO["A-1"] = {"codigo": "A-1", "descricao": valor, "quantidade": Decimal("1")}
        conteudo, _, _ = motor_exportacao.exportar(
            recurso=self.recurso, request=type("R", (), {"user": self.user})(),
            formato=formato, tenant_id=self.tenant.id,
        )
        return conteudo

    def test_csv_neutraliza_os_seis_gatilhos(self):
        for valor in ("=cmd|/c calc!A1", "+1+1", "-1+1", "@SUM(A1)", "\tx", "\rx"):
            texto = self._exportar("csv", valor).decode("utf-8-sig")
            self.assertIn("'" + valor[0], texto, f"gatilho {valor[0]!r} passou cru")

    def test_xlsx_nao_grava_formula_de_verdade(self):
        import io

        from openpyxl import load_workbook

        ws = load_workbook(io.BytesIO(self._exportar("xlsx", "=1+1"))).active
        celula = [c for c in ws[ws.max_row] if c.value == "=1+1"]
        self.assertTrue(celula, "o valor deveria estar lá, como texto")
        self.assertEqual(celula[0].data_type, "s", "R16: saiu como fórmula")

    def test_numero_negativo_nao_e_afetado(self):
        # O escape é só do caminho de TEXTO: número sai tipado (R8) e o `-` de um
        # valor negativo tem de continuar sendo um sinal, não um apóstrofo.
        BANCO["A-1"] = {"codigo": "A-1", "descricao": "Cimento", "quantidade": Decimal("-5")}
        conteudo, _, _ = motor_exportacao.exportar(
            recurso=self.recurso, request=type("R", (), {"user": self.user})(),
            formato="csv", tenant_id=self.tenant.id,
        )
        texto = conteudo.decode("utf-8-sig")
        self.assertIn("-5", texto)
        self.assertNotIn("'-5", texto)


class OrcamentoDeLeituraTests(MotorTestBase):
    """D63 — o parse roda sob teto, e estourar não acusa o arquivo do cliente."""

    def test_estouro_de_memoria_vira_erro_legivel_e_nao_falha_de_leitura(self):
        from unittest.mock import patch

        from funcoesimpexp.limites import Limites

        teto = Limites(
            max_mb=5, formatos=("csv",), memoria_mb=64, segundos=5, plano="free",
        )
        # `MemoryError` levantado DENTRO da leitura é o que o RLIMIT_AS produz.
        # O que se prova aqui é a tradução: o lote fica em ERRO com a mensagem do
        # teto, e não com "arquivo ilegível" — acusar o arquivo por um limite
        # nosso é o defeito que o E2E da PRA.1d pegou na Prancheta.
        with patch("funcoesimpexp.importacao.ler_bytes", side_effect=MemoryError):
            lote = self._preparar([CABECALHO, "A-1;Cimento;10"], teto=teto)
        self.assertEqual(lote.status, Importacao.Status.ERRO)
        self.assertIn("64 MB", lote.mensagem)
        self.assertEqual(lote.linhas.count(), 0, "R1: nada foi materializado")

    def test_sem_teto_declarado_a_leitura_continua_funcionando(self):
        # Chamada interna e caminho de teste não são obrigados a passar orçamento.
        lote = self._preparar([CABECALHO, "A-1;Cimento;10"])
        self.assertEqual(lote.status, Importacao.Status.CONFERINDO)


class EntitlementDeEscritaTests(MotorTestBase):
    """D66 — importar exige o módulo; exportar é portabilidade."""

    def _admin(self, modules=()):
        user = Usuario(username="dono", tenant=self.tenant)
        user.iam_permissions = []
        user.is_tenant_admin = True
        user.iam_modules = list(modules)
        return user

    def test_admin_sem_o_modulo_nao_importa(self):
        from funcoesimpexp.acesso import exigir_acesso_total

        with self.assertRaises(PermissionDenied) as ctx:
            exigir_acesso_total(self._admin(), self.recurso, escrita=True)
        self.assertIn("teste", str(ctx.exception))

    def test_admin_sem_o_modulo_continua_exportando(self):
        from funcoesimpexp.acesso import exigir_acesso_total

        # Portabilidade: cancelar o módulo não pode prender o dado que já existe.
        exigir_acesso_total(self._admin(), self.recurso, escrita=False)

    def test_admin_com_o_modulo_importa(self):
        from funcoesimpexp.acesso import exigir_acesso_total

        exigir_acesso_total(self._admin(modules=["teste"]), self.recurso, escrita=True)

    def test_recurso_de_fundacao_nao_exige_entitlement(self):
        from funcoesimpexp.acesso import exigir_acesso_total

        class Fundacao(type(self.recurso)):
            entitlement = None

        exigir_acesso_total(self._admin(), Fundacao(), escrita=True)


class RastroDeExportacaoTests(MotorTestBase):
    """D67 — a R12 chega à trilha append-only, e não só à linha do tenant."""

    def test_exportar_grava_auditlog_com_recorte_e_volume(self):
        from nucleo.models import AuditLog

        BANCO["A-1"] = {"codigo": "A-1", "descricao": "Cimento", "quantidade": Decimal("1")}
        motor_exportacao.exportar(
            recurso=self.recurso, request=type("R", (), {"user": self.user})(),
            formato="csv", recorte={"codigo": "A-1"}, tenant_id=self.tenant.id,
        )
        log = AuditLog.objects.get(acao=AuditLog.Acao.EXPORTAR_DADOS)
        self.assertEqual(log.detalhes["recurso"], self.recurso.id)
        self.assertEqual(log.detalhes["linhas"], 1)
        self.assertIn("codigo=A-1", log.detalhes["recorte"])


class FaxinaDaEntradaTests(MotorTestBase):
    """D68 — arquivo preso em `aguarda_de_para` tem prazo."""

    @override_settings(MEDIA_ROOT=tempfile.mkdtemp())
    def test_entrada_velha_perde_o_arquivo_e_mantem_a_linha(self):
        from funcoesimpexp.tasks import DIAS_PARA_ARQUIVO_ORFAO, limpar_entradas_orfas

        lote = Importacao.objects.create(
            tenant_id=self.tenant.id, recurso=self.recurso.id,
            arquivo_nome="torto.csv", status=Importacao.Status.ERRO,
            parametros={"aguarda_de_para": True},
        )
        lote.arquivo.save("torto.csv", ContentFile(b"a;b"), save=True)
        Importacao.objects.filter(pk=lote.pk).update(
            criada_em=timezone.now() - timedelta(days=DIAS_PARA_ARQUIVO_ORFAO + 1),
        )

        self.assertEqual(limpar_entradas_orfas(), 1)
        lote.refresh_from_db()
        self.assertFalse(lote.arquivo, "o arquivo tinha de sumir")
        self.assertFalse(lote.parametros["aguarda_de_para"], "a tela não pode oferecer de-para")
        self.assertTrue(Importacao.objects.filter(pk=lote.pk).exists(), "a linha é rastro")

    @override_settings(MEDIA_ROOT=tempfile.mkdtemp())
    def test_entrada_recente_e_preservada(self):
        from funcoesimpexp.tasks import limpar_entradas_orfas

        lote = Importacao.objects.create(
            tenant_id=self.tenant.id, recurso=self.recurso.id, arquivo_nome="novo.csv",
            status=Importacao.Status.ERRO, parametros={"aguarda_de_para": True},
        )
        lote.arquivo.save("novo.csv", ContentFile(b"a;b"), save=True)
        self.assertEqual(limpar_entradas_orfas(), 0)


