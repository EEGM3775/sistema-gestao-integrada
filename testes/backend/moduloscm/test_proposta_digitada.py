# testes/backend/moduloscm/test_proposta_digitada.py
"""A cotação alimentada por quem compra (dono, 16/10/2026).

"Cotação também pode ser alimentada pelo usuário, tanto com importação quanto
digitando, não só recebendo de fornecedor — a rede ainda não está no ar." E o
mapa comparativo sai para a reunião, em PDF e XLSX, uma ou várias cotações.
"""
from __future__ import annotations

import io
from datetime import date
from decimal import Decimal
from unittest.mock import patch

from django.test import TestCase
from rest_framework.test import APITestCase

from catalogo.models import Item
from fornecedores.models import Fornecedor
from moduloerp.models import CentroCusto
from nucleo.models import Tenant, Usuario
from obras.models import Obra

from funcoesimpexp import importacao as motor
from funcoesimpexp.models import LinhaImportacao
from funcoesimpexp.registro import ContextoInvalido, obter
from moduloscm.models import Cotacao, Demanda, ItemDemanda, Proposta

PERMS = ["scm.cotacao.view", "scm.cotacao.edit"]


class Base:
    @classmethod
    def montar(cls, slug):
        cls.tenant = Tenant.objects.create(slug=slug, nome="Construtora Cotação")
        cls.user = Usuario.objects.create(username=f"comprador_{slug}", tenant=cls.tenant)
        cls.menezes = Fornecedor.objects.create(tenant=cls.tenant, razao_social="MENEZES MATERIAIS", email="m@m.ex")
        cls.casa = Fornecedor.objects.create(tenant=cls.tenant, razao_social="CASA DO CONSTRUTOR", email="c@c.ex")
        obra = Obra.objects.create(tenant=cls.tenant, nome_obra="Rede de Drenagem", endereco="Rua A, 1")
        centro = CentroCusto.objects.create(tenant=cls.tenant, nome="Drenagem", tipo=CentroCusto.Tipo.OBRA, obra=obra)
        cls.cimento = Item.objects.create(tenant=cls.tenant, codigo_interno="CIM", descricao="Cimento CP-II", unidade="sc")
        cls.demanda = Demanda.objects.create(tenant=cls.tenant, centro_custo=centro)
        cls.l_cimento = ItemDemanda.objects.create(tenant=cls.tenant, demanda=cls.demanda, produto=cls.cimento,
                                                   quantidade=Decimal("50"), unidade="sc")
        cls.l_areia = ItemDemanda.objects.create(tenant=cls.tenant, demanda=cls.demanda, descricao_servico="Areia média",
                                                 quantidade=Decimal("8"), unidade="m³")
        cls.cotacao = Cotacao.objects.create(tenant=cls.tenant, demanda=cls.demanda)


class PropostaDigitadaTests(Base, APITestCase):
    @classmethod
    def setUpTestData(cls):
        cls.montar("cot-dig")

    def setUp(self):
        self.user.iam_modules = ["scm"]
        self.user.iam_permissions = list(PERMS)
        self.client.force_authenticate(user=self.user)

    def registrar(self, fornecedor, precos, **extra):
        return self.client.post(f"/api/scm/cotacoes/{self.cotacao.pk}/propostas/", {
            "fornecedor": fornecedor.pk,
            "itens": [{"item_demanda": l.pk, "preco_unitario": p} for l, p in precos],
            **extra,
        }, format="json")

    def test_registrar_calcula_total_e_guarda_a_data_em_que_chegou(self):
        resp = self.registrar(self.menezes, [(self.l_cimento, "32,00"), (self.l_areia, "95")],
                              recebida_em="2026-10-10", prazo_entrega_dias=5, observacao="frete incluso")
        self.assertEqual(resp.status_code, 201, resp.content)
        proposta = Proposta.objects.get(pk=resp.data["proposta_id"])
        self.assertEqual(proposta.valor_total, Decimal("2360.00"))  # 50×32 + 8×95
        self.assertEqual((proposta.recebida_em, proposta.prazo_entrega_dias, proposta.observacao),
                         (date(2026, 10, 10), 5, "frete incluso"))
        self.assertEqual(proposta.itens.count(), 2)

    def test_item_sem_preco_fica_de_fora_e_fornecedor_repetido_e_recusado(self):
        resp = self.registrar(self.menezes, [(self.l_cimento, "30"), (self.l_areia, "")])
        self.assertEqual(resp.status_code, 201, resp.content)
        self.assertEqual(Proposta.objects.get(pk=resp.data["proposta_id"]).itens.count(), 1)
        resp = self.registrar(self.menezes, [(self.l_cimento, "29")])
        self.assertEqual(resp.status_code, 400)
        self.assertIn("já tem proposta", resp.data["detail"])

    def test_editar_e_apagar_so_com_a_cotacao_aberta(self):
        pid = self.registrar(self.casa, [(self.l_cimento, "31")]).data["proposta_id"]
        url = f"/api/scm/cotacoes/{self.cotacao.pk}/propostas/{pid}/"
        resp = self.client.put(url, {"fornecedor": self.casa.pk, "itens": [
            {"item_demanda": self.l_cimento.pk, "preco_unitario": "30.50"},
            {"item_demanda": self.l_areia.pk, "preco_unitario": "98"},
        ]}, format="json")
        self.assertEqual(resp.status_code, 200, resp.content)
        self.assertEqual(Decimal(resp.data["valor_total"]), Decimal("2309.00"))

        Cotacao.objects.filter(pk=self.cotacao.pk).update(status=Cotacao.Status.ADJUDICADA)
        self.assertEqual(self.client.delete(url).status_code, 400)
        Cotacao.objects.filter(pk=self.cotacao.pk).update(status=Cotacao.Status.ABERTA)
        self.assertEqual(self.client.delete(url).status_code, 204)
        self.assertFalse(Proposta.objects.filter(pk=pid).exists())

    def test_sem_permissao_de_editar_cotacao_nao_registra(self):
        self.user.iam_permissions = ["scm.cotacao.view"]
        self.assertEqual(self.registrar(self.menezes, [(self.l_cimento, "30")]).status_code, 403)

    def test_mapa_recomenda_o_menor_total_entre_as_que_cotaram_tudo(self):
        from moduloscm.propostas import mapa_da_cotacao

        self.registrar(self.menezes, [(self.l_cimento, "32"), (self.l_areia, "95")])  # 2360
        self.registrar(self.casa, [(self.l_cimento, "30")])  # 1500, mas sem a areia
        mapa = mapa_da_cotacao(self.cotacao)
        menezes = next(c for c in mapa["propostas"] if c["fornecedor"] == "MENEZES MATERIAIS")
        self.assertEqual(mapa["recomendada"], menezes["proposta"])
        self.assertEqual([c["cotados"] for c in mapa["propostas"]], [1, 2])
        cimento = mapa["itens"][0]
        self.assertEqual([p["menor"] for p in cimento["precos"]], [True, False])

    def test_mapa_sai_em_xlsx_e_pdf_uma_ou_varias(self):
        from openpyxl import load_workbook

        self.registrar(self.menezes, [(self.l_cimento, "32"), (self.l_areia, "95")])
        resp = self.client.get(f"/api/scm/cotacoes/{self.cotacao.pk}/mapa/", {"formato": "xlsx"})
        self.assertEqual(resp.status_code, 200, getattr(resp, "content", b"")[:300])
        ws = load_workbook(io.BytesIO(resp.content)).active
        celulas = [c.value for linha in ws.iter_rows() for c in linha]
        self.assertIn("MENEZES MATERIAIS", celulas)
        self.assertIn(2360.0, celulas)

        outra = Cotacao.objects.create(tenant=self.tenant, demanda=self.demanda)
        resp = self.client.get("/api/scm/cotacoes/mapas/", {"ids": f"{self.cotacao.pk},{outra.pk}", "formato": "xlsx"})
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(load_workbook(io.BytesIO(resp.content)).sheetnames,
                         ["Resumo", f"Cotação {self.cotacao.pk}", f"Cotação {outra.pk}"])

        with patch("funcoesimpexp.formatos.escrita.html_para_pdf", return_value=b"%PDF-1.7 mapa") as pdf:
            resp = self.client.get("/api/scm/cotacoes/mapas/", {"ids": f"{self.cotacao.pk},{outra.pk}"})
        self.assertEqual((resp.status_code, resp["Content-Type"]), (200, "application/pdf"))
        html = pdf.call_args.args[0]
        self.assertIn("Mapas comparativos de cotação", html)
        self.assertIn("MENEZES MATERIAIS", html)

    def test_pacote_respeita_o_teto_do_payment(self):
        with patch("payment.planos.SCM_MAPA_TETO_COTACOES", 1):
            resp = self.client.get("/api/scm/cotacoes/mapas/", {"ids": "1,2"})
        self.assertEqual(resp.status_code, 400)


class PlanilhaDaCotacaoTests(Base, TestCase):
    CAB = "Nº da linha (não altere);Código;Item;Und;Quantidade;Preço unitário"

    @classmethod
    def setUpTestData(cls):
        cls.montar("cot-imp")

    def recurso(self, **contexto):
        recurso = obter("scm.proposta")
        contexto = recurso.validar_contexto(contexto, self.user, self.tenant.id)
        return recurso.com_contexto({**contexto, "tenant_id": str(self.tenant.id)})

    def preparar(self, linhas, **contexto):
        return motor.preparar(
            recurso=self.recurso(**contexto), conteudo="\n".join([self.CAB, *linhas]).encode(),
            nome_arquivo="cotacao.csv", tenant_id=self.tenant.id, usuario_nome="comprador",
        )

    def test_a_ida_traz_as_linhas_da_cotacao_sem_preco(self):
        recurso = self.recurso(cotacao=self.cotacao.pk)
        linhas = [recurso.serializar_objeto(o) for o in recurso.queryset(None, {})]
        self.assertEqual([(l["item_demanda"], l["descricao"], l["preco_unitario"]) for l in linhas],
                         [(self.l_cimento.pk, "Cimento CP-II", None), (self.l_areia.pk, "Areia média", None)])

    def test_a_volta_vira_a_proposta_do_fornecedor_e_linha_sem_preco_nao_trava(self):
        lote = self.preparar([f"{self.l_cimento.pk};CIM;Cimento;sc;50;32,00", f"{self.l_areia.pk};;Areia;m³;8;"],
                             cotacao=self.cotacao.pk, fornecedor=self.menezes.pk,
                             recebida_em="2026-10-09", prazo_entrega_dias="7")
        self.assertEqual([l.situacao for l in lote.linhas.order_by("numero")],
                         [LinhaImportacao.Situacao.NOVA, LinhaImportacao.Situacao.NOVA])
        motor.aplicar(lote, recurso=obter("scm.proposta").com_contexto(lote.contexto))
        proposta = Proposta.objects.get(cotacao=self.cotacao, fornecedor=self.menezes)
        self.assertEqual((proposta.valor_total, proposta.recebida_em, proposta.prazo_entrega_dias),
                         (Decimal("1600.00"), date(2026, 10, 9), 7))
        self.assertEqual(proposta.itens.count(), 1)

    def test_linha_de_outra_cotacao_trava_e_fornecedor_com_proposta_e_recusado(self):
        outra = ItemDemanda.objects.create(tenant=self.tenant, demanda=Demanda.objects.create(
            tenant=self.tenant, centro_custo=self.demanda.centro_custo), quantidade=Decimal("1"))
        lote = self.preparar([f"{outra.pk};;Outro;un;1;10"], cotacao=self.cotacao.pk, fornecedor=self.casa.pk)
        self.assertEqual(lote.linhas.get().situacao, LinhaImportacao.Situacao.INVALIDA)

        Proposta.objects.create(tenant=self.tenant, cotacao=self.cotacao, fornecedor=self.casa)
        with self.assertRaises(ContextoInvalido):
            self.recurso(cotacao=self.cotacao.pk, fornecedor=self.casa.pk)
