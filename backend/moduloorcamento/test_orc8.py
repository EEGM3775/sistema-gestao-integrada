"""Testes da ORC.8 — a carga da base oficial num clique.

O que estes testes protegem, em ordem de quanto custa errar:

1. **Preço em branco não é zero.** É a regra 1 do §4.1 e a mais cara de quebrar:
   um insumo gravado a R$ 0,00 atravessa a planilha do cliente sem alarme.
2. **`SD` é onerado e `CD` é desonerado.** A dobra de nomenclatura é da fonte;
   trocar os dois inverte o custo de toda composição com mão de obra.
3. **A analítica é guardada uma vez, não 54.** É a conta que torna a carga
   nacional possível, e é fácil regredir sem que nada quebre — só fica lento e
   gordo.
4. **Classe sem tradução para a carga** (D112 nesta porta).
5. **Só o dono da plataforma carrega.**

O arquivo de teste é construído aqui, célula a célula, e isso é deliberado: o
ZIP oficial tem 16 MB e não pode entrar no repositório (é a lição do D41, onde
as amostras que provaram os motores ficaram fora do git). O que está escrito
abaixo é, ao mesmo tempo, o teste e a **documentação executável** do layout que
a Caixa publica.
"""
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO
from unittest import mock
from zipfile import ZipFile

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase
from django.utils import timezone
from rest_framework.test import APITestCase

from nucleo.frota import MINUTOS_ATE_ORFA
from nucleo.models import Tenant, Usuario
from payment import planos

from . import carga, oficial, services
from .models import (
    AnaliticaNacional,
    BaseReferencial,
    CargaBase,
    ComposicaoReferenciaItem,
    Fonte,
    FonteOficial,
    InsumoReferencia,
)

# As UFs do arquivo sintético. Cinco é o mínimo que `_cabecalho_largo` aceita
# como "isto é formato largo" — usar exatamente o mínimo é de propósito: se
# alguém subir o `MINIMO_DE_UFS` sem pensar, estes testes acusam.
UFS_DO_TESTE = ["AC", "BA", "MG", "RJ", "SP"]


def _planilha_larga() -> bytes:
    """Monta um XLSX no formato nacional (01/2025 em diante).

    Reproduz três coisas que o arquivo real tem e que um mock ingênuo não teria:
    o bloco de título antes do cabeçalho, a **linha de UFs separada** da linha de
    campos, e preço em branco numa UF (o "não houve pesquisa").
    """
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)

    def montar(nome: str, linhas: list[tuple], com_classe: bool):
        aba = wb.create_sheet(nome)
        aba.append(["SISTEMA NACIONAL DE PESQUISA DE CUSTOS E ÍNDICES"])
        aba.append([f"Referência: {nome} — preços medianos"])
        aba.append([])
        # linha das UFs, acima da linha dos campos — como o arquivo real
        aba.append([None, None, None, None] + UFS_DO_TESTE)
        aba.append(
            ["CÓDIGO", "DESCRIÇÃO", "UNIDADE", "CLASSIFICAÇÃO" if com_classe else "ORIGEM"]
            + ["PREÇO"] * len(UFS_DO_TESTE)
        )
        for linha in linhas:
            aba.append(list(linha))

    montar("ISD", [
        # código, descrição, unidade, classe, AC, BA, MG, RJ, SP
        ("1379", "CIMENTO PORTLAND CP II-32", "KG", "MATERIAL",
         "0,71", "0,58", "0,60", "0,63", "0,55"),
        # SERVENTE não tem pesquisa no AC: célula vazia, e isso NÃO é zero
        ("88309", "SERVENTE COM ENCARGOS COMPLEMENTARES", "H", "MÃO DE OBRA",
         "", "19,42", "20,10", "22,05", "23,80"),
        ("00370", "AREIA MEDIA - POSTO JAZIDA", "M3", "MATERIAL",
         "1.234,56", "98,00", "102,30", "115,00", "121,40"),
    ], com_classe=True)
    montar("ICD", [
        ("1379", "CIMENTO PORTLAND CP II-32", "KG", "MATERIAL",
         "0,71", "0,58", "0,60", "0,63", "0,55"),
        ("88309", "SERVENTE COM ENCARGOS COMPLEMENTARES", "H", "MÃO DE OBRA",
         "", "15,90", "16,40", "18,00", "19,20"),
    ], com_classe=True)
    montar("CSD", [
        ("92873", "CONCRETAGEM DE PILARES FCK 25 MPA", "M3", "ESTRUTURAS",
         "512,00", "412,80", "430,00", "455,10", "470,90"),
    ], com_classe=False)
    montar("CCD", [
        ("92873", "CONCRETAGEM DE PILARES FCK 25 MPA", "M3", "ESTRUTURAS",
         "498,00", "398,20", "415,00", "440,00", "455,00"),
    ], com_classe=False)

    analitica = wb.create_sheet("Analítico")
    analitica.append(["COMPOSIÇÕES ANALÍTICAS — COEFICIENTES NACIONAIS"])
    analitica.append([])
    analitica.append([
        "GRUPO", "CODIGO DA COMPOSICAO", "DESCRICAO DA COMPOSICAO",
        "UNIDADE DA COMPOSICAO", "TIPO DE ITEM", "CODIGO DO ITEM",
        "DESCRICAO DO ITEM", "UNIDADE DO ITEM", "COEFICIENTE",
    ])
    for codigo_item, descricao, unidade, coeficiente in [
        ("1379", "CIMENTO PORTLAND CP II-32", "KG", "350,00000"),
        ("88309", "SERVENTE COM ENCARGOS COMPLEMENTARES", "H", "2,50000"),
        ("00370", "AREIA MEDIA - POSTO JAZIDA", "M3", "0,50000"),
    ]:
        analitica.append([
            "ESTRUTURAS", "92873", "CONCRETAGEM DE PILARES FCK 25 MPA", "M3",
            "INSUMO", codigo_item, descricao, unidade, coeficiente,
        ])

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _pacote(nome_planilha: str = "SINAPI_Referencia_2026_07.xlsx") -> bytes:
    """O ZIP como a Caixa publica: a planilha de referência + companhia."""
    buffer = BytesIO()
    with ZipFile(buffer, "w") as pacote:
        pacote.writestr(nome_planilha, _planilha_larga())
        # Os arquivos menores que viajam junto e que o `escolher_planilha`
        # precisa **não** escolher.
        pacote.writestr("SINAPI_maodeobra_2026_07.txt", "irrelevante")
        pacote.writestr("SINAPI_familias.csv", "codigo;familia\n1;X\n")
    return buffer.getvalue()


class LeituraDoFormatoLargoTests(SimpleTestCase):
    """O que o motor entende do arquivo — sem tocar no banco."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.leitura = carga.ler_pacote(_pacote(), "SINAPI-2026-07.zip")

    def test_escolhe_a_planilha_de_referencia_e_nao_a_companhia(self):
        self.assertEqual(self.leitura.planilha, "SINAPI_Referencia_2026_07.xlsx")

    def test_as_quatro_abas_viram_publicacoes_por_uf_e_regime(self):
        publicacoes = self.leitura.publicacoes
        # 5 UFs × 2 regimes (o arquivo do teste não traz ISE/CSE)
        self.assertEqual(len(publicacoes), 10)
        self.assertIn(("SP", BaseReferencial.Regime.ONERADO), publicacoes)
        self.assertIn(("SP", BaseReferencial.Regime.DESONERADO), publicacoes)

    def test_sd_e_onerado_e_cd_e_desonerado(self):
        """A dobra de nomenclatura da fonte, resolvida no lugar certo.

        Trocar estes dois não quebra nada visível: inverte o custo de toda
        composição com mão de obra, calado.
        """
        por_regime = {
            (t.tipo, t.regime): t for t in self.leitura.tabelas
        }
        onerado = por_regime[(InsumoReferencia.Tipo.INSUMO, BaseReferencial.Regime.ONERADO)]
        desonerado = por_regime[
            (InsumoReferencia.Tipo.INSUMO, BaseReferencial.Regime.DESONERADO)
        ]
        servente_onerado = onerado.itens_da_uf("SP")["88309"]
        servente_desonerado = desonerado.itens_da_uf("SP")["88309"]
        # A aba SD (sem desoneração = onerado) é a mais cara: ela carrega a folha
        self.assertEqual(servente_onerado.custo_unitario, Decimal("23.80"))
        self.assertEqual(servente_desonerado.custo_unitario, Decimal("19.20"))
        self.assertGreater(
            servente_onerado.custo_unitario, servente_desonerado.custo_unitario
        )

    def test_preco_em_branco_e_ausencia_e_nao_zero(self):
        """A regra 1 do §4.1, que é a mais cara de quebrar."""
        tabela = next(
            t for t in self.leitura.tabelas
            if t.tipo == InsumoReferencia.Tipo.INSUMO
            and t.regime == BaseReferencial.Regime.ONERADO
        )
        no_acre = tabela.itens_da_uf("AC")
        em_sao_paulo = tabela.itens_da_uf("SP")
        self.assertNotIn("88309", no_acre, "servente sem pesquisa no AC não pode entrar")
        self.assertIn("88309", em_sao_paulo)
        # e o que entrou no AC não virou zero
        self.assertTrue(all(item.custo_unitario > 0 for item in no_acre.values()))

    def test_numero_brasileiro_com_milhar_e_decimal(self):
        tabela = next(
            t for t in self.leitura.tabelas if t.tipo == InsumoReferencia.Tipo.INSUMO
        )
        self.assertEqual(tabela.itens_da_uf("AC")["00370"].custo_unitario, Decimal("1234.56"))

    def test_analitica_e_achada_pelo_cabecalho_e_nao_pelo_nome_da_aba(self):
        self.assertEqual(self.leitura.aba_analitica, "Analítico")
        self.assertEqual(len(self.leitura.analitica), 3)
        cimento = next(
            linha for linha in self.leitura.analitica if linha["codigo_item"] == "1379"
        )
        self.assertEqual(cimento["coeficiente"], Decimal("350.00000"))
        self.assertEqual(cimento["codigo_composicao"], "92873")

    def test_resumo_conta_por_publicacao_antes_de_gravar(self):
        resumo = self.leitura.resumo()
        self.assertEqual(resumo["total_publicacoes"], 10)
        self.assertEqual(resumo["linhas_analiticas"], 3)
        acre_onerado = next(
            p for p in resumo["publicacoes"]
            if p["uf"] == "AC" and p["regime"] == BaseReferencial.Regime.ONERADO
        )
        # 3 insumos na aba, 1 sem pesquisa no AC
        self.assertEqual(acre_onerado["insumos"], 2)
        self.assertEqual(acre_onerado["composicoes"], 1)


class ArquivoRecusadoTests(SimpleTestCase):
    """O que o motor recusa — e recusa dizendo por quê."""

    def test_zip_sem_planilha(self):
        buffer = BytesIO()
        with ZipFile(buffer, "w") as pacote:
            pacote.writestr("leiame.txt", "nada aqui")
        with self.assertRaisesMessage(carga.CargaInvalida, "Nenhuma planilha"):
            carga.ler_pacote(buffer.getvalue(), "x.zip")

    def test_arquivo_que_nao_e_zip_nem_xlsx(self):
        with self.assertRaisesMessage(carga.CargaInvalida, "não é um ZIP"):
            carga.ler_pacote(b"<html>pagina de downloads</html>", "index.html")

    def test_formato_longo_e_recusado_com_o_caminho_certo(self):
        """Publicação anterior a 01/2025 não entra por aqui — e a mensagem diz
        para onde ir, que é o comando com um arquivo por UF."""
        from openpyxl import Workbook

        wb = Workbook()
        aba = wb.active
        aba.append(["CÓDIGO", "DESCRIÇÃO", "UNIDADE", "PREÇO MEDIANO R$"])
        aba.append(["1379", "CIMENTO", "KG", "0,58"])
        buffer = BytesIO()
        wb.save(buffer)
        with self.assertRaisesMessage(carga.CargaInvalida, "importar_base_referencial"):
            carga.ler_pacote(buffer.getvalue(), "sinapi_sp.xlsx")

    def test_planilha_solta_nao_e_aberta_como_pacote(self):
        """Um `.xlsx` **é** um ZIP (OOXML). Perguntar 'é ZIP?' antes de olhar a
        extensão fazia a planilha do plano B ser descompactada em `sheet1.xml`,
        com o erro reclamando que não havia planilha dentro da planilha."""
        dentro = carga.abrir_pacote(_planilha_larga(), "SINAPI_Referencia.xlsx")
        self.assertEqual(list(dentro), ["SINAPI_Referencia.xlsx"])


class GravacaoDaCargaTests(APITestCase):
    """A corrida inteira, com banco."""

    def _carga(self, **campos):
        return CargaBase.objects.create(
            fonte=Fonte.SINAPI, competencia=date(2026, 7, 1),
            origem=CargaBase.Origem.UPLOAD, arquivo_nome="SINAPI-2026-07.zip",
            solicitada_por_nome="teste", **campos,
        )

    def test_carga_grava_as_publicacoes_e_a_analitica_uma_vez_so(self):
        registro = carga.executar(self._carga(), _pacote())
        self.assertEqual(registro.estado, CargaBase.Estado.CONCLUIDA, registro.erro)

        publicacoes = BaseReferencial.objects.filter(
            fonte=Fonte.SINAPI, mes_referencia=date(2026, 7, 1)
        )
        self.assertEqual(publicacoes.count(), 10)

        # **A conta que justifica a `AnaliticaNacional`**: 3 linhas guardadas uma
        # vez, e não 3 × 10 publicações. No arquivo real seriam 67 mil contra 3,6
        # milhões por mês.
        self.assertEqual(AnaliticaNacional.objects.count(), 3)
        self.assertEqual(ComposicaoReferenciaItem.objects.count(), 0)

    def test_publicacao_guarda_o_regime_certo_e_a_procedencia(self):
        carga.executar(self._carga(), _pacote())
        sp_onerado = BaseReferencial.objects.get(
            fonte=Fonte.SINAPI, uf="SP", mes_referencia=date(2026, 7, 1),
            regime_encargos=BaseReferencial.Regime.ONERADO,
        )
        self.assertEqual(sp_onerado.total_insumos, 3)
        self.assertEqual(sp_onerado.total_composicoes, 1)
        self.assertIn("SP", sp_onerado.rotulo)
        servente = InsumoReferencia.objects.get(base=sp_onerado, codigo="88309")
        self.assertEqual(servente.custo_unitario, Decimal("23.8000"))

    def test_uf_sem_pesquisa_nao_ganha_a_linha(self):
        carga.executar(self._carga(), _pacote())
        acre = BaseReferencial.objects.get(
            fonte=Fonte.SINAPI, uf="AC", mes_referencia=date(2026, 7, 1),
            regime_encargos=BaseReferencial.Regime.ONERADO,
        )
        self.assertFalse(
            InsumoReferencia.objects.filter(base=acre, codigo="88309").exists()
        )
        self.assertEqual(acre.total_insumos, 2)

    def test_recarregar_o_mesmo_mes_e_idempotente(self):
        carga.executar(self._carga(), _pacote())
        antes = InsumoReferencia.objects.count()
        segunda = carga.executar(self._carga(), _pacote())
        self.assertEqual(segunda.estado, CargaBase.Estado.CONCLUIDA)
        self.assertEqual(InsumoReferencia.objects.count(), antes)
        self.assertEqual(BaseReferencial.objects.count(), 10)
        self.assertEqual(AnaliticaNacional.objects.count(), 3)

    def test_somente_resumo_nao_grava_nada(self):
        registro = carga.executar(self._carga(somente_resumo=True), _pacote())
        self.assertEqual(registro.estado, CargaBase.Estado.CONCLUIDA)
        self.assertEqual(registro.resumo["total_publicacoes"], 10)
        self.assertEqual(BaseReferencial.objects.count(), 0)
        self.assertEqual(AnaliticaNacional.objects.count(), 0)

    def test_arquivo_invalido_vira_carga_falhou_com_motivo(self):
        registro = carga.executar(self._carga(), b"nao sou um zip")
        self.assertEqual(registro.estado, CargaBase.Estado.FALHOU)
        self.assertIn("não é um ZIP", registro.erro)
        self.assertEqual(BaseReferencial.objects.count(), 0)


class ClasseSemTraducaoNaCargaTests(APITestCase):
    """O D112 nesta porta: a fonte respondeu e nós não entendemos."""

    def setUp(self):
        services.esquecer_classes_declaradas()

    def test_classe_desconhecida_para_a_carga_antes_de_gravar(self):
        with mock.patch.object(
            services, "classificacao_desconhecida",
            side_effect=lambda texto, fonte=None: (texto or "").strip().upper() == "MATERIAL",
        ):
            registro = carga.executar(
                CargaBase.objects.create(
                    fonte=Fonte.SINAPI, competencia=date(2026, 7, 1),
                    origem=CargaBase.Origem.UPLOAD,
                ),
                _pacote(),
            )
        self.assertEqual(registro.estado, CargaBase.Estado.AGUARDANDO)
        self.assertIn("MATERIAL", registro.erro)
        self.assertIn("dados/classes/sinapi.csv", registro.erro)
        # e **nada** foi gravado: publicação meio importada parece completa
        self.assertEqual(BaseReferencial.objects.count(), 0)
        self.assertEqual(AnaliticaNacional.objects.count(), 0)

    def test_aceitar_sem_classe_e_a_resposta_a_essa_pergunta(self):
        with mock.patch.object(
            services, "classificacao_desconhecida",
            side_effect=lambda texto, fonte=None: (texto or "").strip().upper() == "MATERIAL",
        ):
            registro = carga.executar(
                CargaBase.objects.create(
                    fonte=Fonte.SINAPI, competencia=date(2026, 7, 1),
                    origem=CargaBase.Origem.UPLOAD, aceitar_sem_classe=True,
                ),
                _pacote(),
            )
        self.assertEqual(registro.estado, CargaBase.Estado.CONCLUIDA, registro.erro)
        self.assertEqual(BaseReferencial.objects.count(), 10)
        # o registro guarda quais eram, mesmo tendo seguido
        self.assertIn("MATERIAL", registro.resumo["classes_desconhecidas"])

    def test_classe_da_composicao_nao_entra_na_conferencia(self):
        """Em composição o campo traz o capítulo ("ESTRUTURAS"), que não é
        natureza nem pretende ser — reportá-lo encheria a tela de falso alarme."""
        leitura = carga.ler_pacote(_pacote(), "x.zip")
        self.assertNotIn("ESTRUTURAS", leitura.classes)
        self.assertIn("MATERIAL", leitura.classes)


class ResolvedorDaAnaliticaTests(APITestCase):
    """Dois armazenamentos, uma leitura — a razão de `componentes_da_composicao`."""

    def setUp(self):
        carga.executar(
            CargaBase.objects.create(
                fonte=Fonte.SINAPI, competencia=date(2026, 7, 1),
                origem=CargaBase.Origem.UPLOAD,
            ),
            _pacote(),
        )
        self.base = BaseReferencial.objects.get(
            uf="SP", regime_encargos=BaseReferencial.Regime.ONERADO
        )
        self.composicao = InsumoReferencia.objects.get(base=self.base, codigo="92873")

    def test_composicao_resolve_pela_analitica_nacional(self):
        componentes = services.componentes_da_composicao(self.composicao)
        self.assertEqual(len(componentes), 3)
        codigos = {c.filho.codigo for c in componentes}
        self.assertEqual(codigos, {"1379", "88309", "00370"})
        # os filhos são as linhas **daquela publicação** — preço de SP, não do AC
        cimento = next(c for c in componentes if c.filho.codigo == "1379")
        self.assertEqual(cimento.filho.custo_unitario, Decimal("0.5500"))
        self.assertEqual(cimento.coeficiente, Decimal("350.00000"))

    def test_componentes_virtuais_tem_chave_unica_para_a_tela(self):
        componentes = services.componentes_da_composicao(self.composicao)
        chaves = [c.componente_id for c in componentes]
        self.assertEqual(len(set(chaves)), len(chaves))
        self.assertTrue(all(chave < 0 for chave in chaves), "virtual se declara virtual")

    def test_filho_sem_preco_na_uf_fica_de_fora(self):
        """A regra 1 valendo até o fim: uma composição exibida com insumo a
        R$ 0,00 mentiria sobre o próprio custo."""
        acre = BaseReferencial.objects.get(
            uf="AC", regime_encargos=BaseReferencial.Regime.ONERADO
        )
        composicao = InsumoReferencia.objects.get(base=acre, codigo="92873")
        codigos = {c.filho.codigo for c in services.componentes_da_composicao(composicao)}
        self.assertEqual(codigos, {"1379", "00370"})

    def test_analitica_da_propria_publicacao_vence(self):
        """Quem importou pelo comando conferiu aquele arquivo linha a linha; uma
        republicação nacional não muda o que já foi conferido."""
        cimento = InsumoReferencia.objects.get(base=self.base, codigo="1379")
        ComposicaoReferenciaItem.objects.create(
            composicao=self.composicao, filho=cimento, coeficiente=Decimal("1.00000"),
        )
        componentes = services.componentes_da_composicao(self.composicao)
        self.assertEqual(len(componentes), 1)
        self.assertEqual(componentes[0].coeficiente, Decimal("1.00000"))
        self.assertGreater(componentes[0].componente_id, 0)

    def test_endpoint_de_composicao_enxerga_a_analitica_nacional(self):
        tenant = Tenant.objects.create(slug="orc8", nome="Construtora ORC8")
        user = Usuario.objects.create(username="orcamentista8", tenant=tenant)
        user.iam_modules = ["orcamento"]
        user.iam_permissions = ["orcamento.base.view"]
        self.client.force_authenticate(user=user)
        resposta = self.client.get(
            f"/api/orcamento/referencias/{self.composicao.pk}/composicao/"
        )
        self.assertEqual(resposta.status_code, 200)
        self.assertEqual(len(resposta.data["componentes"]), 3)
        primeiro = resposta.data["componentes"][0]
        self.assertIn("custo_no_item", primeiro)
        self.assertIn("filho_descricao", primeiro)


class DownloadTests(SimpleTestCase):
    """A pausa, o teto e a diferença entre 'ainda não saiu' e 'fui cortado'."""

    class _Resposta:
        def __init__(self, status, corpo=b"", headers=None):
            self.status_code = status
            self._corpo = corpo
            self.headers = headers or {}

        def iter_content(self, chunk_size=0):
            for inicio in range(0, len(self._corpo), max(chunk_size, 1)):
                yield self._corpo[inicio:inicio + max(chunk_size, 1)]

        def close(self):
            pass

    def test_404_e_publicacao_ainda_nao_divulgada_e_nao_tenta_de_novo(self):
        sessao = mock.Mock()
        sessao.get.return_value = self._Resposta(404)
        with self.assertRaises(oficial.CargaIndisponivel) as capturado:
            oficial.baixar("https://exemplo/sinapi.zip", sessao=sessao)
        self.assertEqual(capturado.exception.motivo, "nao_publicado")
        self.assertEqual(sessao.get.call_count, 1, "404 não melhora com espera")

    def test_429_insiste_com_backoff_e_aceita_quando_libera(self):
        sessao = mock.Mock()
        sessao.get.side_effect = [
            self._Resposta(429),
            self._Resposta(200, b"conteudo", {"Content-Length": "8"}),
        ]
        with mock.patch("moduloorcamento.oficial.time.sleep") as dormiu:
            baixado = oficial.baixar("https://exemplo/sinapi.zip", sessao=sessao)
        self.assertEqual(baixado.conteudo, b"conteudo")
        self.assertEqual(baixado.tentativas, 2)
        dormiu.assert_called_once_with(planos.ORCAMENTO_CARGA_ESPERA_INICIAL_S)

    def test_desiste_depois_das_tentativas_do_dono_e_propoe_o_plano_b(self):
        sessao = mock.Mock()
        sessao.get.return_value = self._Resposta(429)
        with mock.patch("moduloorcamento.oficial.time.sleep"):
            with self.assertRaises(oficial.CargaIndisponivel) as capturado:
                oficial.baixar("https://exemplo/sinapi.zip", sessao=sessao)
        self.assertEqual(capturado.exception.motivo, "bloqueado")
        self.assertEqual(sessao.get.call_count, planos.ORCAMENTO_CARGA_TENTATIVAS)
        self.assertIn("à mão", str(capturado.exception))

    def test_teto_corta_durante_o_stream_e_nao_depois(self):
        sessao = mock.Mock()
        # sem Content-Length: o teto tem de valer mesmo quando o servidor não diz
        sessao.get.return_value = self._Resposta(200, b"x" * (3 * 1024 * 1024))
        with self.assertRaises(oficial.CargaIndisponivel) as capturado:
            oficial.baixar("https://exemplo/sinapi.zip", sessao=sessao, teto_mb=1)
        self.assertEqual(capturado.exception.motivo, "grande_demais")

    def test_content_length_grande_demais_nem_comeca(self):
        sessao = mock.Mock()
        sessao.get.return_value = self._Resposta(
            200, b"", {"Content-Length": str(400 * 1024 * 1024)}
        )
        with self.assertRaises(oficial.CargaIndisponivel):
            oficial.baixar("https://exemplo/sinapi.zip", sessao=sessao)

    def test_sem_url_diz_que_falta_o_template(self):
        with self.assertRaises(oficial.CargaIndisponivel) as capturado:
            oficial.baixar("")
        self.assertEqual(capturado.exception.motivo, "sem_template")

    def test_competencia_alvo_desconta_a_defasagem_atravessando_o_ano(self):
        self.assertEqual(oficial.competencia_alvo(1, date(2026, 3, 20)), date(2026, 2, 1))
        self.assertEqual(oficial.competencia_alvo(2, date(2026, 1, 5)), date(2025, 11, 1))


class TemplateDaFonteTests(SimpleTestCase):
    """A caixa de texto que evita o deploy."""

    def test_resolve_ano_mes_e_uf(self):
        fonte = FonteOficial(
            fonte=Fonte.SINAPI,
            template_url="https://caixa/{ano}/{mes}/SINAPI_{ano_mes}_{uf}.zip",
        )
        self.assertEqual(
            fonte.url_para(date(2026, 7, 1), uf="sp"),
            "https://caixa/2026/07/SINAPI_202607_SP.zip",
        )

    def test_template_com_chave_desconhecida_devolve_vazio_e_nao_explode(self):
        """URL malformada é problema de quem a escreveu, e a carga sabe dizer
        isso melhor do que um KeyError no log do worker."""
        fonte = FonteOficial(fonte=Fonte.SINAPI, template_url="https://caixa/{trimestre}.zip")
        self.assertEqual(fonte.url_para(date(2026, 7, 1)), "")


class ApiDaCargaTests(APITestCase):
    """Quem pode disparar — e o que a tela recebe."""

    @classmethod
    def setUpTestData(cls):
        cls.tenant = Tenant.objects.create(slug="orc8-api", nome="Construtora")
        cls.dono = Usuario.objects.create(username="dono", tenant=cls.tenant)
        cls.orcamentista = Usuario.objects.create(username="orcamentista", tenant=cls.tenant)

    def _autenticar(self, user, plataforma: bool):
        user.iam_modules = ["orcamento"]
        user.iam_permissions = ["orcamento.base.view", "orcamento.base.adotar"]
        user.is_platform_admin = plataforma
        self.client.force_authenticate(user=user)

    def test_orcamentista_com_o_modulo_nao_carrega_base_mae(self):
        """Ter o módulo e todas as permissões do catálogo não basta: a base-mãe
        vale para todas as empresas."""
        self._autenticar(self.orcamentista, plataforma=False)
        self.assertEqual(self.client.get("/api/orcamento/cargas/").status_code, 403)
        self.assertEqual(self.client.get("/api/orcamento/fontes-oficiais/").status_code, 403)
        resposta = self.client.post(
            "/api/orcamento/cargas/", {"fonte": "sinapi"}, format="json"
        )
        self.assertEqual(resposta.status_code, 403)
        self.assertEqual(CargaBase.objects.count(), 0)

    def test_dono_cadastra_a_fonte_e_a_competencia_sugerida_vem_junto(self):
        self._autenticar(self.dono, plataforma=True)
        resposta = self.client.post("/api/orcamento/fontes-oficiais/", {
            "fonte": "sinapi",
            "template_url": "https://caixa/SINAPI_{ano_mes}.zip",
            "defasagem_meses": 1,
            "dia_do_mes": 12,
        }, format="json")
        self.assertEqual(resposta.status_code, 201, resposta.data)
        self.assertTrue(resposta.data["carga_automatica"])
        self.assertIsNotNone(resposta.data["competencia_sugerida"])

    def test_disparar_enfileira_e_responde_202(self):
        """O `delay` é `on_commit`, e por isso o teste precisa capturar o commit.

        Não é detalhe de teste: enfileirar **antes** do commit deixaria o worker
        buscar uma `CargaBase` que ainda não existe — corrida que só aparece em
        produção, com fila rápida e banco ocupado.
        """
        self._autenticar(self.dono, plataforma=True)
        with mock.patch("moduloorcamento.tasks.carregar_base_oficial.delay") as enfileirou:
            with self.captureOnCommitCallbacks(execute=True):
                resposta = self.client.post(
                    "/api/orcamento/cargas/",
                    {"fonte": "sinapi", "competencia": "2026-07", "somente_resumo": True},
                    format="json",
                )
        self.assertEqual(resposta.status_code, 202, resposta.data)
        registro = CargaBase.objects.get()
        self.assertEqual(registro.competencia, date(2026, 7, 1))
        self.assertEqual(registro.solicitada_por_nome, "dono")
        self.assertTrue(registro.somente_resumo)
        enfileirou.assert_called_once_with(registro.pk)

    def test_duas_cargas_da_mesma_fonte_ao_mesmo_tempo_sao_recusadas(self):
        self._autenticar(self.dono, plataforma=True)
        CargaBase.objects.create(
            fonte=Fonte.SINAPI, competencia=date(2026, 7, 1),
            estado=CargaBase.Estado.BAIXANDO,
        )
        resposta = self.client.post(
            "/api/orcamento/cargas/", {"fonte": "sinapi"}, format="json"
        )
        self.assertEqual(resposta.status_code, 400)
        self.assertIn("andamento", str(resposta.data))

    def test_corrida_orfa_nao_trava_a_fonte_para_sempre(self):
        """Apanhado no E2E: a carga enfileirada para um worker que não estava
        ouvindo ficou `pendente`, e travou a fonte **em definitivo** — nenhuma
        carga nova entrava, e a tela não oferecia saída porque, para ela, havia
        uma rodando.

        É a mesma morte violenta que o `nucleo.liberar_cadeiras_presas` trata na
        frota, e a resposta é a mesma: parou de crescer o carimbo, deixou de
        estar viva.
        """
        self._autenticar(self.dono, plataforma=True)
        presa = CargaBase.objects.create(
            fonte=Fonte.SINAPI, competencia=date(2026, 7, 1),
            estado=CargaBase.Estado.PENDENTE,
        )
        # `atualizada_em` é `auto_now`: só um UPDATE cru o joga para trás.
        CargaBase.objects.filter(pk=presa.pk).update(
            atualizada_em=timezone.now() - timedelta(minutes=MINUTOS_ATE_ORFA + 1)
        )

        with mock.patch("moduloorcamento.tasks.carregar_base_oficial.delay"):
            with self.captureOnCommitCallbacks(execute=True):
                resposta = self.client.post(
                    "/api/orcamento/cargas/", {"fonte": "sinapi"}, format="json"
                )
        self.assertEqual(resposta.status_code, 202, resposta.data)

        # E a órfã não some: vira `falhou` com o motivo escrito, senão o
        # histórico guardaria uma corrida eternamente "na fila".
        presa.refresh_from_db()
        self.assertEqual(presa.estado, CargaBase.Estado.FALHOU)
        self.assertIn("órfã", presa.erro)

    def test_carga_viva_nao_e_confundida_com_orfa(self):
        """O carimbo é reescrito a cada degrau (baixando, lendo, gravando cada
        UF), então uma corrida honesta nunca chega perto do prazo."""
        self._autenticar(self.dono, plataforma=True)
        viva = CargaBase.objects.create(
            fonte=Fonte.SINAPI, competencia=date(2026, 7, 1),
            estado=CargaBase.Estado.GRAVANDO,
        )
        CargaBase.objects.filter(pk=viva.pk).update(
            atualizada_em=timezone.now() - timedelta(minutes=MINUTOS_ATE_ORFA - 1)
        )
        resposta = self.client.post(
            "/api/orcamento/cargas/", {"fonte": "sinapi"}, format="json"
        )
        self.assertEqual(resposta.status_code, 400)
        viva.refresh_from_db()
        self.assertEqual(viva.estado, CargaBase.Estado.GRAVANDO)

    def test_plano_b_aceita_o_arquivo_e_guarda_a_procedencia(self):
        self._autenticar(self.dono, plataforma=True)
        arquivo = SimpleUploadedFile(
            "SINAPI-2026-07.zip", _pacote(), content_type="application/zip"
        )
        with mock.patch("moduloorcamento.tasks.carregar_base_oficial.delay"):
            resposta = self.client.post(
                "/api/orcamento/cargas/enviar/",
                {"fonte": "sinapi", "competencia": "2026-07", "arquivo": arquivo},
                format="multipart",
            )
        self.assertEqual(resposta.status_code, 202, resposta.data)
        registro = CargaBase.objects.get()
        self.assertEqual(registro.origem, CargaBase.Origem.UPLOAD)
        self.assertTrue(registro.arquivo)
        self.assertEqual(registro.arquivo_nome, "SINAPI-2026-07.zip")

    def test_carga_nao_expoe_o_binario(self):
        """O que sustenta a linha `cargas_base: "—"` da tabela de portas do D70.

        O ZIP é entrada operacional: escrito pela API, lido uma vez pelo worker e
        **nunca servido**. Se o campo voltar ao serializer, o DRF publica a URL
        em `/media/` — o caminho que aquele invariante existe para fechar — e
        nada mais quebra para avisar.
        """
        self._autenticar(self.dono, plataforma=True)
        registro = CargaBase.objects.create(
            fonte=Fonte.SINAPI, competencia=date(2026, 7, 1),
            origem=CargaBase.Origem.UPLOAD,
            arquivo=SimpleUploadedFile("x.zip", b"conteudo"),
            arquivo_nome="x.zip",
        )
        resposta = self.client.get(f"/api/orcamento/cargas/{registro.pk}/")
        self.assertEqual(resposta.status_code, 200)
        self.assertNotIn("arquivo", resposta.data)
        self.assertEqual(resposta.data["arquivo_nome"], "x.zip")
        self.assertNotIn("media", str(resposta.data))

    def test_repetir_e_a_resposta_a_carga_que_ficou_aguardando(self):
        self._autenticar(self.dono, plataforma=True)
        parada = CargaBase.objects.create(
            fonte=Fonte.SINAPI, competencia=date(2026, 7, 1),
            estado=CargaBase.Estado.AGUARDANDO, erro="classe sem tradução",
        )
        with mock.patch("moduloorcamento.tasks.carregar_base_oficial.delay"):
            resposta = self.client.post(
                f"/api/orcamento/cargas/{parada.pk}/repetir/",
                {"aceitar_sem_classe": True}, format="json",
            )
        self.assertEqual(resposta.status_code, 202, resposta.data)
        nova = CargaBase.objects.exclude(pk=parada.pk).get()
        self.assertTrue(nova.aceitar_sem_classe)
        self.assertEqual(nova.competencia, parada.competencia)
        # a que parou continua no histórico: nada é apagado
        parada.refresh_from_db()
        self.assertEqual(parada.estado, CargaBase.Estado.AGUARDANDO)

    def test_estado_conta_a_tela_o_que_ela_precisa_para_oferecer_o_botao(self):
        self._autenticar(self.dono, plataforma=True)
        FonteOficial.objects.create(
            fonte=Fonte.SINAPI, template_url="https://caixa/{ano_mes}.zip"
        )
        CargaBase.objects.create(
            fonte=Fonte.SINAPI, competencia=date(2026, 6, 1),
            estado=CargaBase.Estado.CONCLUIDA,
        )
        resposta = self.client.get("/api/orcamento/cargas/estado/")
        self.assertEqual(resposta.status_code, 200)
        self.assertEqual(len(resposta.data["fontes"]), 1)
        self.assertEqual(resposta.data["fontes_com_motor"], [Fonte.SINAPI])
        self.assertIsNone(resposta.data["em_andamento"])
        self.assertIn("sinapi", resposta.data["ultima_por_fonte"])
        self.assertEqual(resposta.data["teto_mb"], planos.ORCAMENTO_CARGA_TETO_MB)


class AgendamentoTests(APITestCase):
    """A varredura que busca — e as três coisas que ela se recusa a fazer."""

    def _rodar(self, dia=12):
        from .tasks import agendar_cargas_do_mes

        with mock.patch("moduloorcamento.tasks.carregar_base_oficial.delay") as enfileirou:
            relatorio = agendar_cargas_do_mes(referencia=f"2026-08-{dia:02d}")
        return relatorio, enfileirou

    def test_dispara_no_dia_declarado_pelo_superuser(self):
        FonteOficial.objects.create(
            fonte=Fonte.SINAPI, template_url="https://caixa/{ano_mes}.zip",
            dia_do_mes=12, defasagem_meses=1,
        )
        relatorio, enfileirou = self._rodar(dia=12)
        self.assertEqual(len(relatorio["disparadas"]), 1)
        self.assertEqual(relatorio["disparadas"][0]["competencia"], "2026-07-01")
        enfileirou.assert_called_once()

    def test_nao_dispara_em_outro_dia(self):
        FonteOficial.objects.create(
            fonte=Fonte.SINAPI, template_url="https://caixa/{ano_mes}.zip", dia_do_mes=12,
        )
        relatorio, enfileirou = self._rodar(dia=13)
        self.assertEqual(relatorio["disparadas"], [])
        enfileirou.assert_not_called()

    def test_fonte_sem_template_nao_e_tentada(self):
        """A decisão nº 1 da tasks.py: um job que inventa fonte é um job que um
        dia importa dado errado com a confiança de um robô."""
        FonteOficial.objects.create(fonte=Fonte.SINAPI, template_url="", dia_do_mes=12)
        relatorio, enfileirou = self._rodar()
        self.assertEqual(relatorio["disparadas"], [])
        self.assertIn("sem template", relatorio["ignoradas"][0]["motivo"])
        enfileirou.assert_not_called()

    def test_fonte_sem_leitor_nao_e_tentada(self):
        FonteOficial.objects.create(
            fonte=Fonte.ORSE, template_url="https://orse/{ano_mes}.zip", dia_do_mes=12,
        )
        relatorio, _ = self._rodar()
        self.assertIn("não lê o formato", relatorio["ignoradas"][0]["motivo"])

    def test_competencia_ja_carregada_nao_roda_de_novo(self):
        FonteOficial.objects.create(
            fonte=Fonte.SINAPI, template_url="https://caixa/{ano_mes}.zip",
            dia_do_mes=12, defasagem_meses=1,
        )
        CargaBase.objects.create(
            fonte=Fonte.SINAPI, competencia=date(2026, 7, 1),
            estado=CargaBase.Estado.CONCLUIDA,
        )
        relatorio, enfileirou = self._rodar()
        self.assertIn("já carregada", relatorio["ignoradas"][0]["motivo"])
        enfileirou.assert_not_called()

    def test_resumo_anterior_nao_conta_como_carregada(self):
        """`somente_resumo` leu e contou, não gravou — deixar isso passar por
        carga faria o mês nunca entrar."""
        FonteOficial.objects.create(
            fonte=Fonte.SINAPI, template_url="https://caixa/{ano_mes}.zip",
            dia_do_mes=12, defasagem_meses=1,
        )
        CargaBase.objects.create(
            fonte=Fonte.SINAPI, competencia=date(2026, 7, 1),
            estado=CargaBase.Estado.CONCLUIDA, somente_resumo=True,
        )
        relatorio, enfileirou = self._rodar()
        self.assertEqual(len(relatorio["disparadas"]), 1)
        enfileirou.assert_called_once()


class CausaDaLinhaSemPrecoNovoTests(APITestCase):
    """A pergunta do dono (21/08/2026): o que acontece com a linha que **não tem
    preço na base mais atual** quando alguém troca a data-base.

    A resposta é: ela fica com a data antiga, entra na prévia como aviso, nunca é
    aplicada, e a grade a marca em vermelho depois (`avisos_de_data_base`). Isso
    é o D99 e continua valendo.

    O que a ORC.8 mudou é que agora há **duas causas** para "não achei este
    código aqui", e elas pedem ações opostas — daí estes testes.
    """

    def setUp(self):
        carga.executar(
            CargaBase.objects.create(
                fonte=Fonte.SINAPI, competencia=date(2026, 7, 1),
                origem=CargaBase.Origem.UPLOAD,
            ),
            _pacote(),
        )
        self.acre = BaseReferencial.objects.get(
            uf="AC", regime_encargos=BaseReferencial.Regime.ONERADO
        )

    def test_sem_pesquisa_na_uf_diz_onde_o_item_existe(self):
        """O servente não tem preço no AC, mas existe no mesmo mês em BA/MG/RJ/SP.

        Dizer "não existe na publicação de destino" seria verdade literal e
        mentira prática: mandaria procurar substituto para um insumo vivo.
        """
        from . import reprecificacao

        faltantes = reprecificacao._onde_mais_existe({"88309"}, self.acre)
        self.assertIn("88309", faltantes)
        self.assertEqual(faltantes["88309"], ["BA", "MG", "RJ", "SP"])

    def test_aposentado_pela_fonte_nao_aparece_em_lugar_nenhum(self):
        """O código que a fonte tirou do catálogo não existe em UF alguma — e é
        esse silêncio que separa as duas causas."""
        from . import reprecificacao

        faltantes = reprecificacao._onde_mais_existe({"99999"}, self.acre)
        self.assertEqual(faltantes, {})

    def test_a_consulta_ignora_a_propria_publicacao_de_destino(self):
        """Senão o cimento, que **tem** preço no AC, apareceria como se também
        estivesse em outro lugar — e a mensagem falaria de um problema que não
        existe."""
        from . import reprecificacao

        faltantes = reprecificacao._onde_mais_existe({"1379"}, self.acre)
        self.assertNotIn("AC", faltantes.get("1379", []))
