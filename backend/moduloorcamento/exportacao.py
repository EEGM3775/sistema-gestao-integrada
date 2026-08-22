# backend/moduloorcamento/exportacao.py
"""ORC.5b — a planilha sai do sistema: XLSX e PDF.

Exportar não é conveniência: é a **entrega**. A planilha vai para o edital, para
o cliente e para o cartório, e nesses destinos o arquivo *é* o produto. Por isso
o que sai daqui carrega o cabeçalho completo (obra, versão, data-base, regime de
encargos, BDI) — uma planilha de custo sem data-base declarada não é conferível,
e é a primeira coisa que um analista de licitação procura.

Duas saídas com propósitos diferentes:

- **XLSX** é para trabalhar: o destinatário vai filtrar, somar e comparar. Sai
  com as colunas separadas e os números como número (não texto), senão a primeira
  coisa que a pessoa faz é reconverter tudo.
- **PDF** é para arquivar e assinar: layout fixo, totais fechados, sem risco de
  alguém "ajustar" uma célula no caminho.

O sintético e o analítico são o mesmo dado em dois níveis de leitura (a
convenção da própria SINAPI) — quem escolhe é quem exporta.
"""
from __future__ import annotations

import io
from decimal import Decimal

from django.template.loader import render_to_string
from django.utils import timezone

from . import precificacao
from .models import ItemOrcado, ItemOrcadoInsumo, Orcamento
from .planilha import dinheiro


class ExportacaoIndisponivel(Exception):
    """Biblioteca de exportação ausente no ambiente."""


CABECALHO_PLANILHA = (
    "Item", "Fonte", "Código", "Descrição", "Und", "Quantidade",
    "Custo unitário", "Total custo direto", "BDI %", "Total com BDI", "Peso %",
)


def _linhas_da_planilha(orcamento: Orcamento, *, analitico: bool) -> list[dict]:
    """Árvore achatada na ordem de leitura da EAP.

    Sintético × analítico é a convenção da **própria SINAPI**, e não um nível de
    zoom na EAP: a peça sintética traz a planilha inteira, um serviço por linha,
    com o custo unitário fechado; a analítica abre, embaixo de cada serviço, os
    insumos que formam aquele custo — coeficiente, preço e a parcela de cada um.

    Antes, o sintético cortava por profundidade de EAP (`1.1.1` sumia), o
    que numa planilha de dois níveis não escondia nada e numa de três escondia
    serviço com preço. Cortar serviço não é resumir: é entregar planilha que não
    fecha com o total declarado no rodapé.
    """
    itens = list(
        ItemOrcado.objects.filter(orcamento=orcamento)
        .select_related("etapa")
        .order_by("codigo_eap", "ordem", "item_orcado_id")
    )

    total = orcamento.valor_total or Decimal("0")
    bdi_padrao = orcamento.bdi_percentual or Decimal("0")
    composicao = precificacao.analitica_por_linha(orcamento) if analitico else {}
    rotulo_categoria = dict(ItemOrcadoInsumo.Categoria.choices)

    linhas: list[dict] = []
    for item in itens:
        nivel = item.codigo_eap.count(".")
        linhas.append({
            "codigo_eap": item.codigo_eap,
            "fonte": (item.fonte_referencia or "").upper(),
            "codigo_referencia": item.codigo_referencia,
            "descricao": item.descricao,
            "unidade": item.unidade,
            "quantidade": item.quantidade if item.custa else None,
            "custo_unitario": item.custo_unitario if item.custa else None,
            "total_custo_direto": item.total_custo_direto,
            "bdi": bdi_padrao,
            "total_com_bdi": item.total_com_bdi,
            "peso": (
                (item.total_com_bdi / total * Decimal("100")).quantize(Decimal("0.01"))
                if total else Decimal("0")
            ),
            "sintetico": item.tipo_linha == ItemOrcado.TipoLinha.SINTETICO,
            "insumo": False,
            "nivel": nivel,
        })

        quantidade = item.quantidade or Decimal("0")
        for insumo in composicao.get(str(item.item_orcado_id), []):
            coeficiente = Decimal(insumo["coeficiente"])
            # A parcela do insumo é multiplicada pela quantidade do serviço de
            # propósito: assim a coluna de custo direto **soma** — os insumos de
            # um serviço fecham no total daquele serviço, e quem confere não
            # precisa refazer a conta à mão. BDI e peso ficam vazios porque a
            # linha de insumo não é uma parcela do preço de venda: ela já está
            # contida no serviço logo acima, e repetir o BDI aqui faria o
            # somatório da coluna dobrar.
            linhas.append({
                "codigo_eap": "",
                "fonte": rotulo_categoria.get(insumo["categoria"], ""),
                "codigo_referencia": "",
                "descricao": f"↳ {insumo['descricao']}",
                "unidade": insumo["unidade"],
                "quantidade": coeficiente * quantidade,
                "custo_unitario": Decimal(insumo["custo_unitario"]),
                "total_custo_direto": Decimal(insumo["custo_no_item"]) * quantidade,
                "bdi": None,
                "total_com_bdi": None,
                "peso": None,
                "sintetico": False,
                "insumo": True,
                "nivel": min(nivel + 1, 3),
            })

    return linhas


def contexto_do_cabecalho(orcamento: Orcamento) -> dict:
    """O cabeçalho que toda peça repete — é o que a torna conferível.

    Planilha de custo sem data-base, regime de encargos e BDI declarados não é
    conferível por quem a recebe, e é a primeira coisa que um analista de
    licitação procura.
    """
    return {
        "obra": orcamento.obra.nome_obra,
        "descricao": orcamento.descricao,
        "versao": orcamento.versao,
        "status": orcamento.get_status_display(),
        "base": (
            orcamento.base_referencial.rotulo or orcamento.base_referencial.rotulo_automatico()
            if orcamento.base_referencial_id else "sem base referencial"
        ),
        "data_base": orcamento.data_base,
        "regime_encargos": orcamento.get_regime_encargos_display(),
        "encargos": orcamento.encargos_sociais_percentual,
        "bdi": orcamento.bdi_percentual,
        "custo_direto_total": orcamento.custo_direto_total,
        "valor_total": orcamento.valor_total,
        "emitido_em": timezone.now(),
    }


# ---------------------------------------------------------------------------
# XLSX
# ---------------------------------------------------------------------------


def planilha_xlsx(orcamento: Orcamento, *, analitico: bool = True) -> bytes:
    """Planilha orçamentária como arquivo Excel, com números como número."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover - dependência de requirements
        raise ExportacaoIndisponivel(
            "openpyxl não instalado — necessário para exportar XLSX."
        ) from exc

    cabecalho = contexto_do_cabecalho(orcamento)
    wb = Workbook()
    ws = wb.active
    ws.title = "Orçamento"

    negrito = Font(bold=True)
    titulo = Font(bold=True, size=13)
    fundo = PatternFill("solid", fgColor="E8EDF5")
    fundo_sintetico = PatternFill("solid", fgColor="F3F5F9")
    borda = Border(bottom=Side(style="thin", color="C9D2E0"))

    ws.append([f"{cabecalho['descricao']} — {cabecalho['obra']}"])
    ws["A1"].font = titulo
    ws.append([
        f"Versão {cabecalho['versao']} ({cabecalho['status']})",
        f"Base: {cabecalho['base']}",
        f"Data-base: {cabecalho['data_base'] or '—'}",
        f"Encargos: {cabecalho['encargos']}% ({cabecalho['regime_encargos']})",
        f"BDI: {cabecalho['bdi']}%",
    ])
    ws.append([f"Emitido em {cabecalho['emitido_em']:%d/%m/%Y %H:%M}"])
    ws.append([])

    linha_cabecalho = ws.max_row + 1
    ws.append(list(CABECALHO_PLANILHA))
    for coluna in range(1, len(CABECALHO_PLANILHA) + 1):
        celula = ws.cell(row=linha_cabecalho, column=coluna)
        celula.font = negrito
        celula.fill = fundo
        celula.border = borda
        celula.alignment = Alignment(horizontal="center", wrap_text=True)

    for linha in _linhas_da_planilha(orcamento, analitico=analitico):
        ws.append([
            linha["codigo_eap"],
            linha["fonte"],
            linha["codigo_referencia"],
            ("    " * linha["nivel"]) + linha["descricao"],
            linha["unidade"],
            float(linha["quantidade"]) if linha["quantidade"] is not None else None,
            float(linha["custo_unitario"]) if linha["custo_unitario"] is not None else None,
            float(linha["total_custo_direto"]),
            float(linha["bdi"]) if linha["bdi"] is not None else None,
            float(linha["total_com_bdi"]) if linha["total_com_bdi"] is not None else None,
            float(linha["peso"]) if linha["peso"] is not None else None,
        ])
        if linha["sintetico"]:
            for coluna in range(1, len(CABECALHO_PLANILHA) + 1):
                celula = ws.cell(row=ws.max_row, column=coluna)
                celula.font = negrito
                celula.fill = fundo_sintetico

    ws.append([])
    ws.append([
        "", "", "", "TOTAL", "", "", "",
        float(cabecalho["custo_direto_total"]), "", float(cabecalho["valor_total"]), "",
    ])
    for coluna in (4, 8, 10):
        ws.cell(row=ws.max_row, column=coluna).font = negrito

    for coluna, largura in enumerate(
        (12, 10, 14, 62, 8, 14, 16, 18, 8, 18, 9), start=1,
    ):
        ws.column_dimensions[get_column_letter(coluna)].width = largura
    for coluna in ("F", "G", "H", "J"):
        for celula in ws[coluna]:
            celula.number_format = "#,##0.00"
    ws.freeze_panes = ws.cell(row=linha_cabecalho + 1, column=1)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def curva_abc_xlsx(orcamento: Orcamento, dados: dict) -> bytes:
    """Curva ABC como planilha — a lista que vai para a negociação de compra."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover
        raise ExportacaoIndisponivel("openpyxl não instalado.") from exc

    wb = Workbook()
    ws = wb.active
    ws.title = "Curva ABC"
    ws.append([f"Curva ABC — {orcamento.descricao} (v{orcamento.versao})"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])

    ws.append(["Classe", "Item", "Descrição", "Und", "Quantidade", "Valor", "%", "% acumulado"])
    for coluna in range(1, 9):
        celula = ws.cell(row=3, column=coluna)
        celula.font = Font(bold=True)
        celula.fill = PatternFill("solid", fgColor="E8EDF5")

    for linha in dados.get("itens", []):
        ws.append([
            linha["classe"], linha["codigo_eap"], linha["descricao"], linha["unidade"],
            float(linha["quantidade"]), float(linha["valor"]),
            float(linha["pct"]), float(linha["pct_acumulado"]),
        ])

    for coluna, largura in enumerate((9, 12, 64, 8, 14, 16, 9, 13), start=1):
        ws.column_dimensions[get_column_letter(coluna)].width = largura

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def cronograma_xlsx(orcamento: Orcamento, dados: dict) -> bytes:
    """Cronograma físico-financeiro: etapas nas linhas, competências nas colunas."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover
        raise ExportacaoIndisponivel("openpyxl não instalado.") from exc

    competencias = dados.get("competencias", [])
    wb = Workbook()
    ws = wb.active
    ws.title = "Cronograma"
    ws.append([f"Cronograma físico-financeiro — {orcamento.descricao} (v{orcamento.versao})"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])

    ws.append(["Etapa", *competencias, "Total"])
    for coluna in range(1, len(competencias) + 3):
        celula = ws.cell(row=3, column=coluna)
        celula.font = Font(bold=True)
        celula.fill = PatternFill("solid", fgColor="E8EDF5")

    for linha in dados.get("linhas", []):
        valores = []
        for competencia in competencias:
            celula = linha["parcelas"].get(competencia)
            valores.append(float(celula["valor_previsto"]) if celula else 0.0)
        ws.append([linha["etapa_nome"], *valores, float(linha["total"])])

    ws.append([])
    totais = []
    for competencia in competencias:
        soma = sum(
            float(l["parcelas"][competencia]["valor_previsto"])
            for l in dados.get("linhas", []) if competencia in l["parcelas"]
        )
        totais.append(soma)
    ws.append(["Desembolso previsto", *totais, float(dados.get("total", 0))])
    for coluna in range(1, len(competencias) + 3):
        ws.cell(row=ws.max_row, column=coluna).font = Font(bold=True)

    ws.column_dimensions["A"].width = 34
    for coluna in range(2, len(competencias) + 3):
        ws.column_dimensions[get_column_letter(coluna)].width = 15

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------


def planilha_pdf(orcamento: Orcamento, *, analitico: bool = True) -> bytes:
    """Planilha em PDF pelo WeasyPrint, a partir de um template HTML.

    HTML → PDF em vez de desenho programático porque a peça precisa de cabeçalho
    repetido por página, quebra que não parte uma linha ao meio e numeração —
    tudo isso o motor de impressão do WeasyPrint resolve com CSS, e uma
    biblioteca de desenho exigiria reimplementar cada um.

    O `write_pdf()` não está aqui de propósito: quem emite é o
    [funcoesimpexp](../funcoesimpexp/formatos/escrita.py), que fecha o buscador
    de URL do WeasyPrint (**D73**). Duas chamadas diretas ao WeasyPrint seriam
    duas políticas de rede, e a segunda nasceria aberta.
    """
    from funcoesimpexp.formatos.escrita import EscritaIndisponivel, html_para_pdf

    html = render_to_string("moduloorcamento/planilha.html", {
        "cabecalho": contexto_do_cabecalho(orcamento),
        "linhas": _linhas_da_planilha(orcamento, analitico=analitico),
        "analitico": analitico,
    })
    try:
        return html_para_pdf(html)
    except EscritaIndisponivel as exc:
        raise ExportacaoIndisponivel(str(exc)) from exc


def nome_do_arquivo(orcamento: Orcamento, *, sufixo: str, extensao: str) -> str:
    """`orcamento-residencial-alfa-v2-analitico.xlsx` — legível na pasta de downloads."""
    base = f"{orcamento.descricao}-v{orcamento.versao}-{sufixo}"
    limpo = "".join(c if c.isalnum() or c in "-_" else "-" for c in base.lower())
    while "--" in limpo:
        limpo = limpo.replace("--", "-")
    return f"{limpo.strip('-')}.{extensao}"


def totais_conferencia(orcamento: Orcamento) -> dict:
    """Números que o rodapé de qualquer exportação repete, para conferência."""
    return {
        "custo_direto": str(dinheiro(orcamento.custo_direto_total or Decimal("0"))),
        "valor_total": str(dinheiro(orcamento.valor_total or Decimal("0"))),
        "linhas": ItemOrcado.objects.filter(orcamento=orcamento).count(),
    }


# ---------------------------------------------------------------------------
# ORC.7 — as peças que acompanham a planilha
# ---------------------------------------------------------------------------
#
# Uma planilha entregue sozinha não é proposta. O envelope de licitação leva
# quatro anexos que até o ORC.6 o módulo não emitia, e cada um responde a uma
# pergunta diferente do analista:
#
#   composição do BDI      "de onde vêm os 24,5%?"          — TCU 2622/2013
#   quadro de encargos     "de onde vêm os 112,8%?"         — metodologia SINAPI
#   relação de insumos     "quanto de aço esta obra usa?"   — e é a lista de compra
#   carta-proposta         a folha assinada com o valor por extenso
#
# Todas saem em XLSX pelo mesmo motor da planilha, e a carta em PDF — ela é peça
# de assinar, e assinar planilha editável não faz sentido.


def _abrir_planilha(titulo: str, subtitulo: str = ""):
    """Workbook novo com o mesmo cabeçalho das outras exportações.

    Existe para as quatro peças novas não repetirem quinze linhas de openpyxl
    cada — e para o cabeçalho continuar igual em todas quando alguém mexer nele.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError as exc:  # pragma: no cover - dependência de requirements
        raise ExportacaoIndisponivel("openpyxl não instalado.") from exc

    wb = Workbook()
    ws = wb.active
    ws.title = titulo[:31]
    ws.append([titulo])
    ws["A1"].font = Font(bold=True, size=13)
    if subtitulo:
        ws.append([subtitulo])
    ws.append([])
    return wb, ws


def _cabecalho_de_colunas(ws, rotulos: tuple, larguras: tuple) -> int:
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    linha = ws.max_row + 1
    ws.append(list(rotulos))
    for coluna in range(1, len(rotulos) + 1):
        celula = ws.cell(row=linha, column=coluna)
        celula.font = Font(bold=True)
        celula.fill = PatternFill("solid", fgColor="E8EDF5")
    for coluna, largura in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(coluna)].width = largura
    return linha


def composicao_bdi_xlsx(orcamento: Orcamento) -> bytes:
    """A **composição analítica do BDI** — parcela a parcela, com a fórmula impressa.

    A fórmula vai no rodapé de propósito: é ela que explica por que os tributos
    dividem em vez de somar, e é a primeira coisa que um analista confere. Sem
    ela a peça vira uma lista de percentuais que não reconstitui o resultado.
    """
    from openpyxl.styles import Font

    from .models import ComposicaoBDI

    composicoes = list(ComposicaoBDI.objects.filter(orcamento=orcamento).order_by("-padrao", "nome"))
    cabecalho = contexto_do_cabecalho(orcamento)
    wb, ws = _abrir_planilha(
        "Composição do BDI",
        f"{cabecalho['descricao']} — {cabecalho['obra']} (v{cabecalho['versao']})",
    )

    if not composicoes:
        ws.append(["Este orçamento usa BDI digitado direto, sem composição analítica."])
        ws.append([f"BDI aplicado: {cabecalho['bdi']}%"])
        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()

    for composicao in composicoes:
        ws.append([f"{composicao.nome}{' (padrão)' if composicao.padrao else ''}"])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        _cabecalho_de_colunas(ws, ("Parcela", "Sigla", "%"), (54, 10, 12))
        for descricao, sigla, valor in (
            ("Administração central", "AC", composicao.administracao_central),
            ("Seguro", "S", composicao.seguro),
            ("Risco", "R", composicao.risco),
            ("Garantia", "G", composicao.garantia),
            ("Despesas financeiras", "DF", composicao.despesas_financeiras),
            ("Lucro", "L", composicao.lucro),
            ("PIS", "I", composicao.pis),
            ("COFINS", "I", composicao.cofins),
            ("ISS", "I", composicao.iss),
            ("CPRB", "I", composicao.cprb),
        ):
            ws.append([descricao, sigla, float(valor)])
        ws.append(["Tributos (I) — soma", "", float(composicao.tributos)])
        ws.append(["BDI resultante", "", float(composicao.percentual_calculado)])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        ws.cell(row=ws.max_row, column=3).font = Font(bold=True)
        ws.append([])

    ws.append(["BDI = [ (1+AC+S+R+G) × (1+DF) × (1+L) ÷ (1−I) ] − 1"])
    ws.append([
        "Os tributos entram dividindo: a alíquota incide sobre o preço de venda, "
        "não sobre o custo. IRPJ e CSLL não compõem o BDI (Acórdãos 325/2007 e "
        "2622/2013-TCU)."
    ])

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def quadro_encargos_xlsx(orcamento: Orcamento) -> bytes:
    """O **quadro de encargos sociais**, grupo a grupo, com a conta do grupo D.

    A coluna "cálculo" é o que separa esta peça de uma tabela de percentuais: nas
    reincidências ela imprime `A (36,80%) × B (48,21%) ÷ 100`, e o analista
    reconstitui o número sem pedir a memória por e-mail.
    """
    from openpyxl.styles import Font

    from . import encargos as encargos_mod
    from .models import ComposicaoEncargos

    quadros = list(
        ComposicaoEncargos.objects.filter(orcamento=orcamento)
        .prefetch_related("rubricas")
        .order_by("-padrao", "tipo_mao_de_obra")
    )
    cabecalho = contexto_do_cabecalho(orcamento)
    wb, ws = _abrir_planilha(
        "Encargos sociais",
        f"{cabecalho['descricao']} — {cabecalho['obra']} (v{cabecalho['versao']})",
    )

    if not quadros:
        ws.append([
            "Este orçamento usa encargos digitados direto, sem quadro analítico."
        ])
        ws.append([f"Encargos aplicados: {cabecalho['encargos']}% ({cabecalho['regime_encargos']})"])
        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()

    for composicao in quadros:
        dados = encargos_mod.quadro(composicao)
        ws.append([
            f"{composicao.nome} — {composicao.get_tipo_mao_de_obra_display()}, "
            f"{composicao.get_regime_display()}{' (padrão)' if composicao.padrao else ''}"
        ])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        _cabecalho_de_colunas(ws, ("Código", "Rubrica", "%", "Cálculo"), (10, 58, 12, 40))

        for grupo in dados["grupos"]:
            ws.append(["", grupo["rotulo"]])
            ws.cell(row=ws.max_row, column=2).font = Font(bold=True)
            for rubrica in grupo["rubricas"]:
                ws.append([
                    rubrica["codigo"], rubrica["descricao"],
                    float(rubrica["percentual"]),
                    rubrica["conta"] if rubrica["derivada"] else "",
                ])
            ws.append(["", f"Subtotal {grupo['grupo']}", float(grupo["subtotal"]), ""])
            ws.cell(row=ws.max_row, column=2).font = Font(bold=True)
            ws.cell(row=ws.max_row, column=3).font = Font(bold=True)

        ws.append(["", "TOTAL DOS ENCARGOS", float(dados["total"]), "A + B + C + D"])
        ws.cell(row=ws.max_row, column=2).font = Font(bold=True)
        ws.cell(row=ws.max_row, column=3).font = Font(bold=True)
        ws.append([])

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def relacao_insumos_xlsx(orcamento: Orcamento, dados: dict) -> bytes:
    """A **relação consolidada de insumos** — a lista de compra da obra.

    O rodapé imprime a cobertura. Uma relação tirada de um orçamento com 40% das
    composições abertas é uma lista de compra pela metade, e quem a recebe sem o
    aviso vai comprar por ela.
    """
    from openpyxl.styles import Font

    wb, ws = _abrir_planilha(
        "Relação de insumos",
        f"{orcamento.descricao} — v{orcamento.versao}",
    )
    _cabecalho_de_colunas(
        ws,
        ("Fonte", "Código", "Insumo", "Und", "Natureza", "Quantidade",
         "Custo unit. médio", "Valor total", "%", "Serviços"),
        (10, 14, 58, 8, 16, 16, 18, 18, 9, 10),
    )

    for linha in dados.get("itens", []):
        ws.append([
            linha["fonte_referencia"], linha["codigo_referencia"], linha["descricao"],
            linha["unidade"], linha["rotulo_categoria"],
            float(linha["quantidade"]), float(linha["custo_unitario_medio"]),
            float(linha["valor"]), float(linha["pct"]), linha["servicos"],
        ])

    ws.append([])
    ws.append(["", "", "TOTAL", "", "", "", "", float(dados.get("total", 0))])
    ws.cell(row=ws.max_row, column=3).font = Font(bold=True)
    ws.cell(row=ws.max_row, column=8).font = Font(bold=True)
    ws.append([])
    ws.append([
        f"Cobertura: {dados.get('cobertura_pct', '0')}% do custo direto está com "
        f"composição explodida ({dados.get('linhas_explodidas', 0)} de "
        f"{dados.get('linhas_analiticas', 0)} linhas). O restante não aparece nesta relação."
    ])

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def curva_abc_insumos_xlsx(orcamento: Orcamento, dados: dict) -> bytes:
    """Curva ABC de insumos — onde negociar preço, não escopo."""
    from openpyxl.styles import Font

    wb, ws = _abrir_planilha(
        "Curva ABC de insumos",
        f"{orcamento.descricao} — v{orcamento.versao}",
    )
    _cabecalho_de_colunas(
        ws,
        ("Classe", "Fonte", "Código", "Insumo", "Und", "Natureza",
         "Quantidade", "Valor", "%", "% acumulado", "Serviços"),
        (9, 10, 14, 54, 8, 16, 16, 18, 9, 13, 10),
    )

    for linha in dados.get("itens", []):
        ws.append([
            linha["classe"], linha["fonte_referencia"], linha["codigo_referencia"],
            linha["descricao"], linha["unidade"], linha["categoria"],
            float(linha["quantidade"]), float(linha["valor"]),
            float(linha["pct"]), float(linha["pct_acumulado"]), linha["servicos"],
        ])

    ws.append([])
    resumo = dados.get("resumo", {})
    for classe in ("A", "B", "C"):
        bloco = resumo.get(classe, {})
        ws.append([
            classe, "", "", f"{bloco.get('itens', 0)} insumos", "", "", "",
            float(bloco.get("valor", 0)), float(bloco.get("pct", 0)),
        ])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    ws.append([])
    ws.append([
        f"Cobertura: {dados.get('cobertura_pct', '0')}% do custo direto explodido."
    ])

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# Carta-proposta
# ---------------------------------------------------------------------------

_UNIDADES = (
    "", "um", "dois", "três", "quatro", "cinco", "seis", "sete", "oito", "nove",
    "dez", "onze", "doze", "treze", "quatorze", "quinze", "dezesseis",
    "dezessete", "dezoito", "dezenove",
)
_DEZENAS = (
    "", "", "vinte", "trinta", "quarenta", "cinquenta", "sessenta", "setenta",
    "oitenta", "noventa",
)
_CENTENAS = (
    "", "cento", "duzentos", "trezentos", "quatrocentos", "quinhentos",
    "seiscentos", "setecentos", "oitocentos", "novecentos",
)


def _ate_999(numero: int) -> str:
    if numero == 100:
        return "cem"
    centena, resto = divmod(numero, 100)
    dezena, unidade = divmod(resto, 10)
    partes = []
    if centena:
        partes.append(_CENTENAS[centena])
    if resto < 20 and resto:
        partes.append(_UNIDADES[resto])
    else:
        if dezena:
            partes.append(_DEZENAS[dezena])
        if unidade:
            partes.append(_UNIDADES[unidade])
    return " e ".join(partes)


def valor_por_extenso(valor: Decimal) -> str:
    """O valor em palavras — exigência de carta-proposta, não enfeite.

    Edital pede o valor por extenso porque é o que prevalece sobre o algarismo em
    caso de divergência. Escrever à mão numa proposta gerada por sistema é como
    entra o erro de uma casa decimal que ninguém percebe até a abertura.
    """
    valor = Decimal(valor).quantize(Decimal("0.01"))
    inteiro = int(valor)
    centavos = int((valor - inteiro) * 100)

    if inteiro == 0:
        texto = "zero real"
    else:
        blocos: list[tuple[str, int]] = []
        escalas = (
            (1_000_000_000, "bilhão", "bilhões"),
            (1_000_000, "milhão", "milhões"),
            (1_000, "mil", "mil"),
        )
        restante = inteiro
        for base, singular, plural in escalas:
            quantidade, restante = divmod(restante, base)
            if quantidade:
                nome = singular if quantidade == 1 else plural
                prefixo = "" if base == 1_000 and quantidade == 1 else f"{_ate_999(quantidade)} "
                blocos.append((f"{prefixo}{nome}".strip(), base))
        if restante:
            blocos.append((_ate_999(restante), restante))

        # A conjunção antes do último bloco não é estilo: "mil e duzentos" e
        # "mil duzentos" são valores diferentes na leitura de um cheque. A regra
        # é o "e" só quando o último bloco é menor que cem **ou** centena
        # redonda — 2.500 vira "dois mil e quinhentos", 1.234 vira "mil
        # duzentos e trinta e quatro".
        texto = blocos[0][0]
        for posicao, (parte, valor) in enumerate(blocos[1:], start=1):
            ultimo = posicao == len(blocos) - 1
            conjuncao = ultimo and (valor < 100 or (valor < 1000 and valor % 100 == 0))
            texto += (" e " if conjuncao else " ") + parte

        # "um milhão de reais", não "um milhão reais": com escala redonda a
        # preposição é obrigatória.
        if inteiro == 1:
            texto += " real"
        elif inteiro >= 1_000_000 and inteiro % 1_000_000 == 0:
            texto += " de reais"
        else:
            texto += " reais"

    if centavos:
        moeda = "centavo" if centavos == 1 else "centavos"
        texto += f" e {_ate_999(centavos)} {moeda}"
    return texto


def carta_proposta_pdf(orcamento: Orcamento, *, dados: dict | None = None) -> bytes:
    """A folha de rosto que acompanha a planilha — em PDF, porque é para assinar.

    Sai pelo mesmo caminho da planilha (`html_para_pdf` do funcoesimpexp, que
    fecha o buscador de URL do WeasyPrint — D73). Duas chamadas diretas ao
    WeasyPrint seriam duas políticas de rede, e a segunda nasceria aberta.
    """
    from funcoesimpexp.formatos.escrita import EscritaIndisponivel, html_para_pdf

    cabecalho = contexto_do_cabecalho(orcamento)
    total = orcamento.valor_total or Decimal("0")
    contexto = {
        "cabecalho": cabecalho,
        "total": dinheiro(total),
        "total_extenso": valor_por_extenso(total),
        "custo_direto": dinheiro(orcamento.custo_direto_total or Decimal("0")),
        "prazo_dias": (dados or {}).get("prazo_dias"),
        "validade_dias": (dados or {}).get("validade_dias"),
        "proponente": (dados or {}).get("proponente", ""),
        "responsavel": (dados or {}).get("responsavel", ""),
        "crea": (dados or {}).get("crea", ""),
        "cidade": (dados or {}).get("cidade", ""),
        "emitido_em": timezone.localtime(),
    }
    html = render_to_string("moduloorcamento/carta_proposta.html", contexto)
    try:
        return html_para_pdf(html)
    except EscritaIndisponivel as exc:
        raise ExportacaoIndisponivel(str(exc)) from exc
