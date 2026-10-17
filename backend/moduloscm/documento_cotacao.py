# backend/moduloscm/documento_cotacao.py
"""O mapa comparativo da cotação sai para a reunião: PDF e XLSX (dono, 16/10/2026).

- **PDF** em paisagem para projetar ou imprimir: itens nas linhas, um fornecedor
  por coluna, o menor preço de cada linha em destaque, total, prazo, data em que
  o preço chegou, a recomendação e o campo de aprovação. Pela única porta de PDF
  do sistema (`html_para_pdf`, buscador sem rede).
- **XLSX** com o mesmo mapa e os números como número, para quem quer mexer.

Várias cotações de uma vez saem num arquivo só: uma capa com o resumo e um mapa
por cotação (no XLSX, uma aba por cotação). O dado é o de `propostas.
mapa_da_cotacao` — o mesmo comparativo da tela.
"""
from __future__ import annotations

import io

from django.template.loader import render_to_string
from django.utils import timezone

from .documento_pedido import _brl, _qtd
from .propostas import mapa_da_cotacao

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class ExportacaoIndisponivel(RuntimeError):
    pass


def _empresa(tenant_id) -> str:
    from nucleo.models import Tenant

    tenant = Tenant.objects.filter(pk=tenant_id).only("nome").first()
    return tenant.nome if tenant else ""


def _resumo(mapa: dict) -> dict:
    """Uma linha da capa: quem é o recomendado (ou o vencedor) e por quanto."""
    por_id = {c["proposta"]: c for c in mapa["propostas"]}
    escolhida = por_id.get(mapa["vencedora"]) or por_id.get(mapa["recomendada"])
    return {
        "cotacao": mapa["cotacao"], "demanda": mapa["demanda"], "obra": mapa["obra"] or mapa["centro_custo"],
        "status": mapa["status"], "itens": len(mapa["itens"]), "propostas": len(mapa["propostas"]),
        "fornecedor": escolhida["fornecedor"] if escolhida else "",
        "total": escolhida["total"] if escolhida else None,
        "vencedora": bool(mapa["vencedora"]),
    }


def _para_o_template(mapa: dict) -> dict:
    """Valores já formatados (o template não formata dinheiro em pt-BR sozinho)."""
    return {
        **mapa,
        "itens": [{
            **linha, "quantidade_fmt": _qtd(linha["quantidade"]),
            "precos": [{**p, "fmt": _brl(p["valor"]) if p["valor"] is not None else "—"} for p in linha["precos"]],
        } for linha in mapa["itens"]],
        "propostas": [{
            **c, "total_fmt": _brl(c["total"]) if c["total"] is not None else "—",
            "situacao": _situacao(mapa, c),
        } for c in mapa["propostas"]],
        "compacto": len(mapa["propostas"]) > 5,
    }


def pdf_dos_mapas(cotacoes) -> bytes:
    from funcoesimpexp.formatos.escrita import html_para_pdf

    mapas = [mapa_da_cotacao(c) for c in cotacoes]
    html = render_to_string("moduloscm/mapa_cotacao.html", {
        "empresa": _empresa(cotacoes[0].tenant_id),
        "gerado_em": timezone.localtime(),
        "mapas": [_para_o_template(m) for m in mapas],
        "capa": [{**r, "total_fmt": _brl(r["total"]) if r["total"] is not None else "—"} for r in map(_resumo, mapas)]
        if len(mapas) > 1 else None,
    })
    return html_para_pdf(html)


def xlsx_dos_mapas(cotacoes) -> bytes:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover
        raise ExportacaoIndisponivel("openpyxl não instalado — necessário para exportar XLSX.") from exc

    negrito = Font(bold=True)
    cabecalho = PatternFill("solid", fgColor="E8EDF5")
    menor = PatternFill("solid", fgColor="DCFCE7")
    dinheiro = '"R$" #,##0.00'
    mapas = [mapa_da_cotacao(c) for c in cotacoes]
    wb = Workbook()
    primeira = wb.active

    if len(mapas) > 1:
        ws = primeira
        ws.title = "Resumo"
        ws.append([f"Mapas comparativos — {_empresa(cotacoes[0].tenant_id)}"])
        ws["A1"].font = Font(bold=True, size=13)
        ws.append([f"Gerado em {timezone.localtime():%d/%m/%Y %H:%M}"])
        ws.append([])
        ws.append(["Cotação", "Demanda", "Obra / centro de custo", "Situação", "Itens", "Propostas", "Escolhido", "", "Total"])
        for celula in ws[4]:
            celula.font, celula.fill = negrito, cabecalho
        for r in map(_resumo, mapas):
            ws.append([r["cotacao"], r["demanda"], r["obra"], r["status"], r["itens"], r["propostas"], r["fornecedor"],
                       "vencedor" if r["vencedora"] else ("menor total" if r["fornecedor"] else ""),
                       float(r["total"]) if r["total"] is not None else None])
            ws.cell(ws.max_row, 9).number_format = dinheiro
        for col, largura in zip("ABCDEFGHI", (10, 10, 32, 14, 8, 10, 36, 14, 16)):
            ws.column_dimensions[col].width = largura

    for indice, mapa in enumerate(mapas):
        ws = primeira if (indice == 0 and len(mapas) == 1) else wb.create_sheet()
        ws.title = f"Cotação {mapa['cotacao']}"[:31]
        ws.append([f"Mapa comparativo — Cotação #{mapa['cotacao']} · Demanda #{mapa['demanda']}"])
        ws["A1"].font = Font(bold=True, size=13)
        ws.append([" · ".join(x for x in (mapa["obra"], mapa["centro_custo"], mapa["status"]) if x)])
        ws.append([])
        n = len(mapa["propostas"])
        topo = ["Código", "Item", "Und", "Qtd."] + [c["fornecedor"] for c in mapa["propostas"]]
        ws.append(topo)
        linha_topo = ws.max_row
        for celula in ws[linha_topo]:
            celula.font, celula.fill = negrito, cabecalho
            celula.alignment = Alignment(wrap_text=True, vertical="top")
        for item in mapa["itens"]:
            ws.append([item["codigo"], item["descricao"], item["unidade"], float(item["quantidade"])]
                      + [float(p["valor"]) if p["valor"] is not None else None for p in item["precos"]])
            for j, p in enumerate(item["precos"]):
                celula = ws.cell(ws.max_row, 5 + j)
                celula.number_format = dinheiro
                if p["menor"]:
                    celula.fill, celula.font = menor, negrito
        rodape = (
            ("Total", [float(c["total"]) if c["total"] is not None else None for c in mapa["propostas"]], dinheiro),
            ("Itens cotados", [f"{c['cotados']} de {len(mapa['itens'])}" for c in mapa["propostas"]], None),
            ("Prazo (dias)", [c["prazo_entrega_dias"] for c in mapa["propostas"]], None),
            ("Recebida em", [c["recebida_em"] for c in mapa["propostas"]], "DD/MM/YYYY"),
            ("Situação", [_situacao(mapa, c) for c in mapa["propostas"]], None),
        )
        for rotulo, valores, formato in rodape:
            ws.append(["", rotulo, "", ""] + valores)
            ws.cell(ws.max_row, 2).font = negrito
            if formato:
                for j in range(n):
                    ws.cell(ws.max_row, 5 + j).number_format = formato
        if mapa["nenhuma_completa"]:
            ws.append([])
            ws.append(["", "Nenhuma proposta cotou todos os itens: o menor total compara listas diferentes."])
        ws.column_dimensions["A"].width = 12
        ws.column_dimensions["B"].width = 42
        ws.column_dimensions["C"].width = 7
        ws.column_dimensions["D"].width = 10
        for j in range(n):
            ws.column_dimensions[get_column_letter(5 + j)].width = 18
        ws.freeze_panes = ws.cell(linha_topo + 1, 5)

    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue()


def _situacao(mapa: dict, coluna: dict) -> str:
    if coluna["proposta"] == mapa["vencedora"]:
        return "Vencedora"
    if coluna["proposta"] == mapa["recomendada"] and not mapa["vencedora"]:
        return "Menor total"
    return coluna["status_rotulo"]


def nome_do_arquivo(cotacoes, extensao: str) -> str:
    if len(cotacoes) == 1:
        return f"mapa-cotacao-{cotacoes[0].pk}.{extensao}"
    return f"mapas-de-cotacao-{timezone.localdate():%Y%m%d}.{extensao}"
